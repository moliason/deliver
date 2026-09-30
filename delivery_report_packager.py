from __future__ import annotations

import argparse
import copy
import hashlib
import html
import json
import os
import posixpath
import re
import zipfile
from datetime import datetime
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET
from urllib.parse import urlencode, urljoin
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image


def norm(value) -> str:
    # Excel 表头和序号里可能有普通空格，比较前统一去掉。
    return str(value or "").replace(" ", "").strip()


def reply_text(reply: dict, analysis_key: str = "reason_analysis", analysis_label: str = "原因分析") -> str:
    # 文档要求写回格式固定为“原因分类 + 原因分析”。
    return f"原因分类：{reply.get('reason_category', '')}\n{analysis_label}：{reply.get(analysis_key, '')}"


def redmine_config(config: dict) -> dict:
    # 地址和字段由调用方配置，凭据只从环境变量读取。
    redmine = dict(config.get("redmine", {}))
    redmine["url"] = redmine.get("url") or os.getenv("REDMINE_URL", "")
    if not redmine["url"]:
        raise ValueError("缺少 Redmine URL")
    redmine.setdefault("project_id", "")
    redmine.setdefault("status", "Verified")
    redmine.setdefault(
        "custom_fields",
        {
            "match_value": "序号",
            "reason_category": "(研发）原因分类",
            "reason_analysis": "（研发）原因分析1",
            "owner": "整改人",
            "deliverables": "交付物清单",
            "design_commit_id": "架构设计git commit id",
            "main_commit_id": "main&libopt git commit id",
            "pin_commit_id": "pin_function git commit id",
            "reason_analysis2": "（研发）原因分析2",
            "reason_analysis3": "（研发）原因分析3",
        },
    )
    redmine["api_key"] = os.getenv("REDMINE_API_KEY", "")
    if not redmine["api_key"]:
        raise ValueError("缺少环境变量 REDMINE_API_KEY")
    return redmine


def redmine_json(redmine: dict, path: str, params: dict | None = None) -> dict:
    # 调 Redmine REST API，当前只用 issue 列表、issue 详情和附件下载。
    query = f"?{urlencode(params)}" if params else ""
    request = Request(urljoin(redmine["url"].rstrip("/") + "/", path) + query)
    request.add_header("X-Redmine-API-Key", redmine["api_key"])
    with urlopen(request, timeout=30) as response:
        return json.loads(response.read().decode("utf-8"))


def redmine_download(redmine: dict, url: str, target: Path) -> None:
    # 下载 feature 附件到 output/redmine_reports 下，后面直接修改下载后的 xlsx。
    request = Request(url)
    request.add_header("X-Redmine-API-Key", redmine["api_key"])
    target.parent.mkdir(parents=True, exist_ok=True)
    with urlopen(request, timeout=30) as response:
        target.write_bytes(response.read())


def custom_field(issue: dict, name: str) -> str:
    # 子 bug 的原因分类、原因分析、整改人、匹配值都来自自定义字段。
    for field in issue.get("custom_fields", []):
        if field.get("name") == name:
            return str(field.get("value") or "")
    return ""


def issue_fields(issue: dict, fields: dict) -> dict:
    return {
        "id": issue.get("id", ""),
        "subject": issue.get("subject", ""),
        "status": issue.get("status", {}).get("name", ""),
        "priority": issue.get("priority", {}).get("name", ""),
        "assigned_to": issue.get("assigned_to", {}).get("name", ""),
        "category": issue.get("category", {}).get("name", ""),
        "start_date": issue.get("start_date", ""),
        "due_date": issue.get("due_date", ""),
        "done_ratio": issue.get("done_ratio", ""),
        "estimated_hours": issue.get("estimated_hours", ""),
        "reason_category": custom_field(issue, fields.get("reason_category", "")),
        "reason_analysis1": custom_field(issue, fields.get("reason_analysis", "")),
        "reason_analysis2": custom_field(issue, fields.get("reason_analysis2", "")),
        "reason_analysis3": custom_field(issue, fields.get("reason_analysis3", "")),
        "deliverables": custom_field(issue, fields.get("deliverables", "")),
        "design_commit_id": custom_field(issue, fields.get("design_commit_id", "")),
        "main_commit_id": custom_field(issue, fields.get("main_commit_id", "")),
        "pin_commit_id": custom_field(issue, fields.get("pin_commit_id", "")),
    }


def issue_reply(issue: dict, fields: dict, default_match_value: str = "") -> dict:
    match_value = custom_field(issue, fields["match_value"]) or default_match_value or issue.get("subject", "")
    return {
        "match_value": match_value,
        "bug_id": f"#{issue['id']}",
        "bug_title": issue.get("subject", ""),
        "description": issue.get("description") or "",
        "reason_category": custom_field(issue, fields["reason_category"]),
        "reason_analysis": custom_field(issue, fields["reason_analysis"]),
        "reason_analysis2": custom_field(issue, fields.get("reason_analysis2", "")),
        "reason_analysis3": custom_field(issue, fields.get("reason_analysis3", "")),
        "owner": custom_field(issue, fields["owner"]) or issue.get("assigned_to", {}).get("name", ""),
    }


def bug_reply(issue: dict, fields: dict) -> dict:
    return issue_reply(issue, fields)


def has_reason_content(reply: dict) -> bool:
    return bool(reply.get("reason_category") or reply.get("reason_analysis"))


def image_fingerprint(data: bytes) -> str:
    image = Image.open(BytesIO(data)).convert("L").resize((8, 8))
    pixels = list(image.getdata())
    avg = sum(pixels) / len(pixels)
    return "".join("1" if pixel >= avg else "0" for pixel in pixels)


def nonconformity_replies(redmine: dict, feature: dict, child_replies: list[dict]) -> list[dict]:
    # 反馈单也可能有多条整改项，优先按子 bug/关联问题逐条填写；没有子项时才用 feature 自身字段兜底。
    if child_replies:
        return child_replies

    feature_reply = issue_reply(feature, redmine["custom_fields"], "1")
    if has_reason_content(feature_reply):
        return [feature_reply]
    return []


def child_bug_replies(redmine: dict, feature: dict) -> list[dict]:
    # 报告回复来源：优先取父子 bug，同时兼容 Redmine “关联的问题”。
    fields = redmine["custom_fields"]
    feature_id = feature["id"]
    data = redmine_json(redmine, "issues.json", {"parent_id": feature_id, "status_id": "*", "limit": 100})
    issues = []
    for issue in data.get("issues", []):
        if "description" not in issue or "attachments" not in issue:
            issue = redmine_json(redmine, f"issues/{issue['id']}.json", {"include": "attachments"})["issue"]
        issues.append(issue)

    seen_ids = {f"#{issue['id']}" for issue in issues}
    for relation in feature.get("relations", []):
        related_id = relation["issue_to_id"] if relation["issue_id"] == feature_id else relation["issue_id"]
        if f"#{related_id}" in seen_ids:
            continue
        issue = redmine_json(redmine, f"issues/{related_id}.json", {"include": "attachments"})["issue"]
        issues.append(issue)
        seen_ids.add(f"#{related_id}")

    replies = []
    for issue in issues:
        reply = bug_reply(issue, fields)
        reply["image_hashes"] = []
        reply["image_fingerprints"] = []
        for attachment in issue.get("attachments", []):
            if Path(attachment.get("filename", "")).suffix.lower() not in {".png", ".jpg", ".jpeg"}:
                continue
            request = Request(attachment["content_url"])
            request.add_header("X-Redmine-API-Key", redmine["api_key"])
            with urlopen(request, timeout=30) as response:
                data = response.read()
                reply["image_hashes"].append(hashlib.sha256(data).hexdigest())
                reply["image_fingerprints"].append(image_fingerprint(data))
        replies.append(reply)
    return replies


def verified_report_feature_ids(redmine: dict) -> list[int]:
    # 未指定 feature_ids 时，从配置项目里查找符合状态的 issue。
    data = redmine_json(
        redmine,
        "issues.json",
        {"project_id": redmine["project_id"], "status_id": "*", "limit": 100},
    )
    ids = []
    for issue in data.get("issues", []):
        if issue.get("status", {}).get("name") == redmine["status"]:
            ids.append(issue["id"])
    return ids


def report_type(filename: str) -> str:
    if "集成测试" in filename or "系统测试" in filename:
        return "test_report"
    if "开发环节评审表" in filename or "评审表" in filename:
        return "review"
    if "不合格反馈单" in filename or "反馈单" in filename:
        return "nonconformity"
    return ""


def has_test_report_deliverable(deliverables) -> bool:
    return any(report_type(deliverable) == "test_report" for deliverable in deliverables)


def download_redmine_reports(config: dict, output_dir: Path) -> dict:
    # 只下载报告附件；代码文件不在这里处理。
    redmine = redmine_config(config)
    reports = {"review": [], "nonconformity": [], "test_report": []}
    download_dir = Path(redmine.get("download_dir") or output_dir / "redmine_reports")

    feature_ids = redmine.get("feature_ids") or verified_report_feature_ids(redmine)
    for feature_id in feature_ids:
        feature = redmine_json(redmine, f"issues/{feature_id}.json", {"include": "attachments,relations"})["issue"]
        replies = child_bug_replies(redmine, feature)
        for attachment in feature.get("attachments", []):
            filename = attachment["filename"]
            current_report_type = report_type(filename)
            if not current_report_type:
                continue

            path = download_dir / str(feature_id) / filename
            redmine_download(redmine, attachment["content_url"], path)
            reports[current_report_type].append({"path": str(path), "replies": replies})
    return reports


def download_and_fill_feature_reports(
    config: dict,
    feature_id: int,
    download_dir: Path,
    report_types: set[str] | None = None,
) -> list[str]:
    redmine = redmine_config(config)
    feature = redmine_json(redmine, f"issues/{feature_id}.json", {"include": "attachments,relations"})["issue"]
    print(f"FEATURE_FIELDS={json.dumps(issue_fields(feature, redmine['custom_fields']), ensure_ascii=False)}")
    replies = child_bug_replies(redmine, feature)
    changed = []

    for attachment in feature.get("attachments", []):
        current_report_type = report_type(attachment["filename"])
        if not current_report_type or report_types is not None and current_report_type not in report_types:
            continue

        path = Path(download_dir) / attachment["filename"]
        redmine_download(redmine, attachment["content_url"], path)
        item = {"path": str(path), "replies": replies}
        if current_report_type == "review":
            ai_feedback_changed = fill_ai_feedback_review(item)
            if ai_feedback_changed is None:
                changed.extend(fill_review(item))
            else:
                changed.extend(ai_feedback_changed)
        elif current_report_type == "nonconformity":
            item["replies"] = nonconformity_replies(redmine, feature, replies)
            changed.extend(fill_nonconformity(item))
            for reply in item["replies"]:
                print(f"BUG_DOCUMENT={path}")
                print(f"BUG_ITEM={reply['bug_id']} {reply['bug_title']}")
        elif current_report_type == "test_report":
            changed.extend(fill_test_report(item))

    return changed


def find_review_columns(sheet) -> dict:
    # RCU 评审表示例中：指摘序号列在“指摘内容”表头下，回复列为“责任人判断/对应”。
    for row in range(1, 81):
        values = {norm(sheet.cell(row, col).value): col for col in range(1, 31)}
        if norm("指摘内容") in values and norm("责任人判断/对应") in values and norm("指摘者确认") in values:
            no_col = values[norm("指摘内容")]
            reply_col = values[norm("责任人判断/对应")]
            return {
                "header_row": row,
                "no_col": no_col,
                "content_col": no_col + 1,
                "reply_col": reply_col,
                "confirm_col": values[norm("指摘者确认")],
            }
    raise ValueError(f"{sheet.title} 找不到评审表指摘区域")


def find_review_reply_row(sheet, cols: dict, match_values) -> tuple[int | None, float]:
    rows = range(cols["header_row"] + 1, sheet.max_row + 1)
    best_row = None
    best_similarity = 0.0
    for match_value in match_values:
        normalized_match = "".join(str(match_value or "").split()).casefold()
        if not normalized_match:
            continue
        if normalized_match.isdigit():
            row = next(
                (
                    row
                    for row in rows
                    if norm(sheet.cell(row, cols["no_col"]).value) == normalized_match
                ),
                None,
            )
            if row is not None:
                return row, 1.0
            continue

        for row in rows:
            content = "".join(str(sheet.cell(row, cols["content_col"]).value or "").split()).casefold()
            if not content:
                continue
            if normalized_match in content or content in normalized_match:
                return row, 1.0
            similarity = SequenceMatcher(None, normalized_match, content).ratio()
            if similarity > best_similarity:
                best_row = row
                best_similarity = similarity

    if best_similarity >= 0.99:
        return best_row, best_similarity
    return None, best_similarity


def find_review_reply_row_by_image(path: Path, sheet, reply: dict) -> int | None:
    # 纯图片指摘没有文字可匹配时，用子 bug 图片附件和评审表内嵌图片做精确匹配。
    if not reply.get("image_hashes") and not reply.get("image_fingerprints"):
        return None

    exact_rows = []
    similar_rows = []
    for image in getattr(sheet, "_images", []):
        image.ref.seek(0)
        data = image.ref.read()
        image.ref.seek(0)
        row = image.anchor._from.row + 1
        if hashlib.sha256(data).hexdigest() in reply.get("image_hashes", []):
            exact_rows.append(row)
            continue
        fingerprint = image_fingerprint(data)
        if any(
            sum(left != right for left, right in zip(fingerprint, target)) <= 8
            for target in reply.get("image_fingerprints", [])
        ):
            similar_rows.append(row)

    exact_rows = sorted(set(exact_rows))
    if len(exact_rows) == 1:
        return exact_rows[0]
    similar_rows = sorted(set(similar_rows))
    return similar_rows[0] if len(similar_rows) == 1 else None


def worksheet_xml_path(path: Path, sheet) -> str:
    workbook_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relation_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    with zipfile.ZipFile(path, "r") as workbook:
        workbook_xml = ET.fromstring(workbook.read("xl/workbook.xml"))
        workbook_rels = ET.fromstring(workbook.read("xl/_rels/workbook.xml.rels"))
    sheet_rel_id = None
    for item in workbook_xml.findall(f"{{{workbook_namespace}}}sheets/{{{workbook_namespace}}}sheet"):
        if item.attrib.get("name") == sheet.title:
            sheet_rel_id = item.attrib.get(f"{{{relation_namespace}}}id")
            break
    if not sheet_rel_id:
        raise ValueError(f"{sheet.title} 找不到对应 worksheet 关系")
    for item in workbook_rels.findall(f"{{{package_namespace}}}Relationship"):
        if item.attrib.get("Id") == sheet_rel_id:
            target = item.attrib["Target"].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise ValueError(f"{sheet.title} 找不到对应 worksheet XML")


def write_cells_without_repacking_images(
    path: Path,
    sheet,
    values: dict[str, str],
    copied_rows: list[tuple[int, int, str]] | None = None,
    cell_styles: dict[str, int] | None = None,
    column_widths: dict[int, float] | None = None,
) -> None:
    # openpyxl 重新保存带图片的报告时可能丢图；直接改 xlsx 内部 XML，保留原始图片文件。
    xml_namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    markup_namespace = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    drawing_namespace = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
    x14_namespace = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
    x14ac_namespace = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"
    preferred_namespace_prefixes = {
        xml_namespace: "",
        markup_namespace: "mc",
        office_relationship_namespace: "r",
        drawing_namespace: "xdr",
        x14_namespace: "x14",
        x14ac_namespace: "x14ac",
    }
    required_namespace_uris = {
        "x14": x14_namespace,
        "x14ac": x14ac_namespace,
    }
    sheet_xml = worksheet_xml_path(path, sheet)
    temp_path = path.with_suffix(path.suffix + ".tmp")
    copied_rows = copied_rows or []
    cell_styles = cell_styles or {}
    column_widths = column_widths or {}

    def parse_xml(data: bytes):
        namespaces = {}
        for _, (prefix, uri) in ET.iterparse(BytesIO(data), events=("start-ns",)):
            prefix = preferred_namespace_prefixes.get(uri, prefix or "")
            if prefix != "xml" and not re.fullmatch(r"ns\d+", prefix):
                namespaces.setdefault(prefix, uri)

        root = ET.fromstring(data)
        for element in root.iter():
            for name, value in element.attrib.items():
                if name.rsplit("}", 1)[-1] not in {"Ignorable", "Requires"}:
                    continue
                for prefix in value.split():
                    if prefix in required_namespace_uris:
                        namespaces[prefix] = required_namespace_uris[prefix]

        for prefix, uri in namespaces.items():
            ET.register_namespace(prefix, uri)
        return root, namespaces

    def serialize_xml(root, namespaces: dict[str, str]) -> bytes:
        data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        declaration_end = data.find(b"?>")
        root_start = data.find(b"<", declaration_end + 2 if declaration_end >= 0 else 0)
        root_end = data.find(b">", root_start)
        root_tag = data[root_start:root_end + 1].decode("utf-8")
        declared_prefixes = {
            match.group(1) or ""
            for match in re.finditer(r'\sxmlns(?::([A-Za-z_][\w.-]*))?="[^"]+"', root_tag)
        }
        missing_declarations = []
        for prefix, uri in namespaces.items():
            if prefix in declared_prefixes:
                continue
            name = f"xmlns:{prefix}" if prefix else "xmlns"
            missing_declarations.append(f' {name}="{html.escape(uri, quote=True)}"')

        insertion = "".join(missing_declarations).encode("utf-8")
        insert_at = root_end - 1 if data[root_end - 1:root_end] == b"/" else root_end
        return data[:insert_at] + insertion + data[insert_at:]

    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        drawing_paths = set()
        if copied_rows:
            sheet_rels = posixpath.join(
                posixpath.dirname(sheet_xml),
                "_rels",
                f"{posixpath.basename(sheet_xml)}.rels",
            )
            if sheet_rels in source.namelist():
                relationships = ET.fromstring(source.read(sheet_rels))
                for relationship in relationships.findall(f"{{{relationship_namespace}}}Relationship"):
                    if not relationship.attrib.get("Type", "").endswith("/drawing"):
                        continue
                    drawing_target = relationship.attrib["Target"]
                    drawing_paths.add(
                        drawing_target.lstrip("/")
                        if drawing_target.startswith("/")
                        else posixpath.normpath(posixpath.join(posixpath.dirname(sheet_xml), drawing_target))
                    )

        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == sheet_xml:
                root, namespaces = parse_xml(data)
                sheet_data = root.find(f"{{{xml_namespace}}}sheetData")
                merge_cells = root.find(f"{{{xml_namespace}}}mergeCells")

                def set_inline_string(row, coordinate: str, value: str) -> None:
                    cell = next(
                        (
                            current_cell
                            for current_cell in row.findall(f"{{{xml_namespace}}}c")
                            if current_cell.attrib.get("r") == coordinate
                        ),
                        None,
                    )
                    if cell is None:
                        cell = ET.SubElement(row, f"{{{xml_namespace}}}c", {"r": coordinate})
                    for child in list(cell):
                        cell.remove(child)
                    cell.attrib["t"] = "inlineStr"
                    inline = ET.SubElement(cell, f"{{{xml_namespace}}}is")
                    text = ET.SubElement(inline, f"{{{xml_namespace}}}t")
                    text.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
                    text.text = value

                copies_by_row = {}
                for source_row, reply_col, reply_value in copied_rows:
                    copies_by_row.setdefault(source_row, []).append((reply_col, reply_value))

                def inserted_before(row_number: int) -> int:
                    return sum(len(replies) for source_row, replies in copies_by_row.items() if source_row < row_number)

                if copies_by_row:
                    def set_row_number(row, row_number: int) -> None:
                        row.attrib["r"] = str(row_number)
                        for cell in row.findall(f"{{{xml_namespace}}}c"):
                            column = re.match(r"[A-Z]+", cell.attrib.get("r", ""))
                            if column:
                                cell.attrib["r"] = f"{column.group()}{row_number}"

                    original_rows = sorted(
                        sheet_data.findall(f"{{{xml_namespace}}}row"),
                        key=lambda row: int(row.attrib.get("r", "0")),
                    )
                    final_rows = []
                    for row in original_rows:
                        original_row = int(row.attrib.get("r", "0"))
                        final_row = original_row + inserted_before(original_row)
                        set_row_number(row, final_row)
                        final_rows.append(row)
                        for offset, (reply_col, reply_value) in enumerate(copies_by_row.get(original_row, []), start=1):
                            copied_row = copy.deepcopy(row)
                            set_row_number(copied_row, final_row + offset)
                            for cell in copied_row.findall(f"{{{xml_namespace}}}c"):
                                # 新增行只沿用原行格式，不重复指摘内容、指摘者和客户确认文字。
                                for child in list(cell):
                                    cell.remove(child)
                                cell.attrib.pop("t", None)
                            set_inline_string(
                                copied_row,
                                f"{get_column_letter(reply_col)}{final_row + offset}",
                                reply_value,
                            )
                            final_rows.append(copied_row)

                    for row in original_rows:
                        sheet_data.remove(row)
                    for row in final_rows:
                        sheet_data.append(row)

                    if merge_cells is not None:
                        duplicate_refs = []
                        for merged in merge_cells.findall(f"{{{xml_namespace}}}mergeCell"):
                            match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", merged.attrib.get("ref", ""))
                            if not match:
                                continue
                            first_col, first_row, last_col, last_row = match.groups()
                            first_row = int(first_row)
                            last_row = int(last_row)
                            final_first_row = first_row + inserted_before(first_row)
                            final_last_row = last_row + inserted_before(last_row)
                            merged.attrib["ref"] = f"{first_col}{final_first_row}:{last_col}{final_last_row}"
                            if first_row == last_row:
                                duplicate_refs.extend(
                                    f"{first_col}{final_first_row + offset}:{last_col}{final_first_row + offset}"
                                    for offset in range(1, len(copies_by_row.get(first_row, [])) + 1)
                                )
                        for merged_ref in duplicate_refs:
                            ET.SubElement(merge_cells, f"{{{xml_namespace}}}mergeCell", {"ref": merged_ref})
                        merge_cells.attrib["count"] = str(len(merge_cells))

                rows = {row.attrib.get("r"): row for row in sheet_data.findall(f"{{{xml_namespace}}}row")}
                for coordinate, value in values.items():
                    original_row = int(re.search(r"\d+$", coordinate).group())
                    row_no = str(original_row + sum(1 for source_row, _, _ in copied_rows if source_row < original_row))
                    row = rows.get(row_no)
                    if row is None:
                        row = ET.SubElement(sheet_data, f"{{{xml_namespace}}}row", {"r": row_no})
                        rows[row_no] = row
                    column = re.match(r"[A-Z]+", coordinate).group()
                    set_inline_string(row, f"{column}{row_no}", value)

                if cell_styles:
                    cells = {
                        cell.attrib.get("r"): cell
                        for row in rows.values()
                        for cell in row.findall(f"{{{xml_namespace}}}c")
                    }
                    for coordinate, style_id in cell_styles.items():
                        cell = cells.get(coordinate)
                        if cell is None:
                            row_no = re.search(r"\d+$", coordinate).group()
                            cell = ET.SubElement(rows[row_no], f"{{{xml_namespace}}}c", {"r": coordinate})
                        if not style_id:
                            cell.attrib.pop("s", None)
                        else:
                            cell.attrib["s"] = str(style_id)

                if column_widths:
                    columns = root.find(f"{{{xml_namespace}}}cols")
                    if columns is None:
                        columns = ET.Element(f"{{{xml_namespace}}}cols")
                        root.insert(list(root).index(sheet_data), columns)
                    for column_index, width in column_widths.items():
                        ET.SubElement(
                            columns,
                            f"{{{xml_namespace}}}col",
                            {
                                "min": str(column_index),
                                "max": str(column_index),
                                "width": str(width),
                                "customWidth": "1",
                            },
                        )

                dimension = root.find(f"{{{xml_namespace}}}dimension")
                if dimension is not None and rows:
                    dimension_match = re.fullmatch(r"([A-Z]+\d+):([A-Z]+)\d+", dimension.attrib.get("ref", ""))
                    if dimension_match:
                        last_column = dimension_match.group(2)
                        if column_widths:
                            last_column = get_column_letter(
                                max(column_index_from_string(last_column), max(column_widths))
                            )
                        dimension.attrib["ref"] = (
                            f"{dimension_match.group(1)}:{last_column}{max(map(int, rows))}"
                        )
                data = serialize_xml(root, namespaces)
            elif copied_rows and item.filename in drawing_paths:
                root, namespaces = parse_xml(data)
                for row in root.findall(f".//{{{drawing_namespace}}}row"):
                    original_row = int(row.text or "0")
                    row.text = str(original_row + sum(1 for source_row, _, _ in copied_rows if original_row >= source_row))
                data = serialize_xml(root, namespaces)
            target.writestr(item, data)

    os.replace(temp_path, path)


def fill_ai_feedback_review(item: dict) -> list[str] | None:
    path = Path(item["path"])
    workbook = load_workbook(path)
    if "AI反馈结果" not in workbook.sheetnames:
        return None

    sheet = workbook["AI反馈结果"]
    header_row = 2
    headers = {norm(cell.value): cell.column for cell in sheet[header_row]}
    position_header = norm("位置（函数 / 大致行号）")
    if position_header not in headers or "建议" not in headers:
        raise ValueError(f"{sheet.title} 找不到位置（函数 / 大致行号）、建议表头")

    position_col = headers[position_header]
    suggestion_col = headers["建议"]
    reply_col = suggestion_col + 1
    confirm_col = reply_col + 1
    reply_header = norm(sheet.cell(header_row, reply_col).value)
    confirm_header = norm(sheet.cell(header_row, confirm_col).value)
    if reply_header not in {"", norm("责任人判断/对应")} or confirm_header not in {"", norm("指摘者确认")}:
        raise ValueError(f"{sheet.title} 的建议列右侧已有其他内容，不能创建评审回复列")

    creating_columns = not reply_header and not confirm_header
    updates = {
        f"{get_column_letter(reply_col)}{header_row}": "责任人判断/对应",
        f"{get_column_letter(confirm_col)}{header_row}": "指摘者确认",
    }
    position_rows = [
        (
            row,
            "".join(str(sheet.cell(row, position_col).value or "").split()).casefold(),
        )
        for row in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row, position_col).value
    ]
    changed = []
    copied_rows = []
    for reply in item["replies"]:
        # AI 反馈 Bug 的标题对应位置内容，描述通常只有截图。
        normalized_position = "".join(str(reply.get("bug_title") or "").split()).casefold()
        row = next(
            (
                current_row
                for current_row, position in position_rows
                if normalized_position and (normalized_position in position or position in normalized_position)
            ),
            None,
        )
        if row is None:
            print(
                f"AI_REVIEW_REPLY_UNMATCHED={path.name}\t{reply.get('bug_id', '')}\t"
                f"{reply.get('bug_title', '')}"
            )
            continue
        value = reply_text(reply)
        confirmation = norm(sheet.cell(row, confirm_col).value)
        if confirmation and "已确认" not in confirmation:
            copied_rows.append((row, reply_col, value))
        else:
            updates[f"{get_column_letter(reply_col)}{row}"] = value
        changed.append(
            f"{path.name}:AI反馈结果:{reply.get('bug_id') or reply.get('bug_title', '')}"
        )

    cell_styles = {}
    column_widths = {}
    if creating_columns:
        internal_style_col = max(1, suggestion_col - 1)
        for row in range(header_row, sheet.max_row + 1):
            suggestion_coordinate = f"{get_column_letter(suggestion_col)}{row}"
            reply_coordinate = f"{get_column_letter(reply_col)}{row}"
            confirm_coordinate = f"{get_column_letter(confirm_col)}{row}"
            internal_style_id = sheet.cell(row, internal_style_col).style_id
            cell_styles[suggestion_coordinate] = internal_style_id
            cell_styles[reply_coordinate] = internal_style_id
            cell_styles[confirm_coordinate] = sheet.cell(row, suggestion_col).style_id
        suggestion_width = sheet.column_dimensions[get_column_letter(suggestion_col)].width or 30
        column_widths = {reply_col: suggestion_width, confirm_col: 16}

    write_cells_without_repacking_images(
        path,
        sheet,
        updates,
        copied_rows=copied_rows,
        cell_styles=cell_styles,
        column_widths=column_widths,
    )
    return changed


def fill_review(item: dict) -> list[str]:
    # 按子 bug 的匹配值找到指摘序号，写入“责任人判断/对应”。
    path = Path(item["path"])
    workbook = load_workbook(path)
    sheet = workbook.active
    cols = find_review_columns(sheet)
    changed = []
    updates = {}
    copied_rows = []

    for reply in item["replies"]:
        match_value = reply["match_value"]
        if not match_value:
            print(f"REVIEW_REPLY_UNMATCHED={path.name}\t{reply.get('bug_id', '')}\t{reply.get('bug_title', '')}")
            continue
        description = html.unescape(re.sub(r"<[^>]+>|![^!\n]+!|!\[[^\]]*\]\([^)]*\)", " ", reply.get("description") or ""))
        row, similarity = find_review_reply_row(sheet, cols, [match_value, reply.get("bug_title", ""), description])
        if row is None:
            row = find_review_reply_row_by_image(path, sheet, reply)
            if row is not None:
                similarity = 1.0
        if row is None:
            print(
                f"REVIEW_REPLY_UNMATCHED={path.name}\t{reply.get('bug_id', '')}\t"
                f"{reply.get('bug_title', '')}\tbest_similarity={similarity:.4f}"
            )
            continue
        value = reply_text(reply)
        confirmation = norm(sheet.cell(row, cols["confirm_col"]).value)
        if confirmation and norm("已确认") not in confirmation:
            copied_rows.append((row, cols["reply_col"], value))
        else:
            updates[sheet.cell(row, cols["reply_col"]).coordinate] = value
        changed.append(f"{path.name}:{reply['match_value']}")

    if updates or copied_rows:
        write_cells_without_repacking_images(path, sheet, updates, copied_rows)
    return changed


def fill_nonconformity(item: dict) -> list[str]:
    # 不合格反馈单固定结构：A 列序号、H 列整改情况、J 列整改人。
    path = Path(item["path"])
    workbook = load_workbook(path)
    sheet = workbook.active
    changed = []
    data_rows = [r for r in range(5, sheet.max_row + 1) if norm(sheet.cell(r, 1).value)]

    for reply in item["replies"]:
        match_value = norm(reply["match_value"])
        if not match_value:
            continue
        row = next((r for r in data_rows if norm(sheet.cell(r, 1).value) == match_value), None)
        if row is None and len(data_rows) == 1:
            row = data_rows[0]
        if row is None:
            continue
        sheet.cell(row, 8).value = reply_text(reply)
        sheet.cell(row, 10).value = reply.get("owner", "")
        changed.append(f"{path.name}:{reply['match_value']}")

    workbook.save(path)
    return changed


def find_test_report_columns(workbook) -> dict:
    for sheet in workbook.worksheets:
        for row in range(1, min(sheet.max_row, 80) + 1):
            values = {norm(sheet.cell(row, col).value): col for col in range(1, min(sheet.max_column, 30) + 1)}
            if norm("序号") in values and norm("问题") in values and norm("修改方式") in values:
                return {
                    "sheet": sheet,
                    "header_row": row,
                    "no_col": values[norm("序号")],
                    "problem_col": values[norm("问题")],
                    "reply_col": values[norm("修改方式")],
                }
    raise ValueError("找不到测试报告问题区域")


def fill_test_report(item: dict) -> list[str]:
    # 集成测试报告、系统测试报告按序号写入“修改方式”列。
    path = Path(item["path"])
    workbook = load_workbook(path)
    cols = find_test_report_columns(workbook)
    sheet = cols["sheet"]
    changed = []
    updates = {}

    for reply in item["replies"]:
        match_values = [reply.get("match_value", ""), reply.get("bug_title", "")]
        description = html.unescape(re.sub(r"<[^>]+>|![^!\n]+!|!\[[^\]]*\]\([^)]*\)", " ", reply.get("description") or ""))
        if description:
            match_values.append(description)
        if not any(norm(value) for value in match_values) and not (
            reply.get("image_hashes") or reply.get("image_fingerprints")
        ):
            continue
        row = None
        best_similarity = 0.0
        for match_value in match_values:
            normalized_match = "".join(str(match_value or "").split()).casefold()
            if not normalized_match:
                continue
            if normalized_match.isdigit():
                row = next(
                    (
                        r for r in range(cols["header_row"] + 1, sheet.max_row + 1)
                        if norm(sheet.cell(r, cols["no_col"]).value) == normalized_match
                    ),
                    None,
                )
                if row is not None:
                    break
                continue
            for r in range(cols["header_row"] + 1, sheet.max_row + 1):
                content = "".join(str(sheet.cell(r, cols["problem_col"]).value or "").split()).casefold()
                if not content:
                    continue
                if normalized_match in content or content in normalized_match:
                    row = r
                    break
                similarity = SequenceMatcher(None, normalized_match, content).ratio()
                if similarity > best_similarity:
                    best_similarity = similarity
                    if similarity >= 0.99:
                        row = r
            if row is not None:
                break
        if row is None:
            row = find_review_reply_row_by_image(path, sheet, reply)
        if row is None:
            print(
                f"TEST_REPORT_REPLY_UNMATCHED={path.name}\t{reply.get('bug_id', '')}\t"
                f"{reply.get('bug_title', '')}\tbest_similarity={best_similarity:.4f}"
            )
            continue
        updates[sheet.cell(row, cols["reply_col"]).coordinate] = reply_text(reply)
        changed.append(f"{path.name}:{reply['match_value']}")

    if updates:
        write_cells_without_repacking_images(path, sheet, updates)
    return changed


# ---------- Jenkins HTML 报告 ----------
# 输入：build/deliver_output.txt 和 build/package 下的实际归档产物。
# 输出：build/html_report/index.html，用于展示产物、缺失文件原因和关键日志。

def html_report_read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def html_report_parse_log(log_text: str) -> dict:
    missing_items = []
    deliver_items = []
    report_filled = []
    deliver_lists = []
    error_lines = []

    for line in log_text.splitlines():
        line = line.lstrip("\ufeff")
        if line.startswith("MISSING_DELIVER_FILE="):
            parts = line.removeprefix("MISSING_DELIVER_FILE=").split("\t")
            while len(parts) < 4:
                parts.append("")
            missing_items.append({
                "project": parts[0],
                "mode": parts[1],
                "file_type": parts[2],
                "reason": parts[3],
            })
        elif line.startswith("DELIVER_ITEM="):
            parts = line.removeprefix("DELIVER_ITEM=").split("\t")
            while len(parts) < 3:
                parts.append("")
            deliver_items.append({"project": parts[0], "mode": parts[1], "files": parts[2]})
        elif line.startswith("REPORT_FILLED="):
            parts = line.removeprefix("REPORT_FILLED=").split("\t")
            while len(parts) < 2:
                parts.append("")
            report_filled.append({"issue": parts[0], "count": parts[1]})
        elif line.startswith("DELIVER_LIST="):
            deliver_lists.append(line.removeprefix("DELIVER_LIST="))
        elif line.startswith(("Error:", "交付失败:", "归档失败:")) or "Traceback (most recent call last)" in line:
            error_lines.append(line)

    return {
        "missing_items": missing_items,
        "deliver_items": deliver_items,
        "report_filled": report_filled,
        "deliver_lists": deliver_lists,
        "error_lines": error_lines,
    }


def html_report_package_files(package_dir: Path) -> list[dict]:
    if not package_dir.exists():
        return []

    files = []
    for path in sorted(package_dir.iterdir(), key=lambda item: item.name.casefold()):
        if path.is_file():
            files.append({"name": path.name, "size": path.stat().st_size})
    return files


def html_report_format_size(size: int) -> str:
    if size < 1024:
        return f"{size} B"
    if size < 1024 * 1024:
        return f"{size / 1024:.1f} KiB"
    return f"{size / 1024 / 1024:.1f} MiB"


def html_report_table(headers: list[str], rows: list[list[str]]) -> str:
    head = "".join(f"<th>{html.escape(header)}</th>" for header in headers)
    body = []
    for row in rows:
        body.append("<tr>" + "".join(f"<td>{html.escape(str(cell))}</td>" for cell in row) + "</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def html_report_style(status: str) -> str:
    status_color = "#b42318" if status != "成功" else "#16794c"
    return f"""
    body {{ margin: 0; font-family: Arial, "Microsoft YaHei", sans-serif; color: #1f2933; background: #f6f8fb; }}
    header {{ padding: 24px 32px; background: #26364d; color: #fff; }}
    h1 {{ margin: 0 0 8px; font-size: 24px; }}
    main {{ padding: 24px 32px; }}
    section {{ margin-bottom: 24px; background: #fff; border: 1px solid #d9e2ec; border-radius: 6px; overflow: hidden; }}
    h2 {{ margin: 0; padding: 14px 16px; font-size: 17px; background: #eef3f8; border-bottom: 1px solid #d9e2ec; }}
    .summary {{ display: flex; gap: 12px; flex-wrap: wrap; margin-bottom: 20px; }}
    .card {{ background: #fff; border: 1px solid #d9e2ec; border-radius: 6px; padding: 14px 16px; min-width: 160px; }}
    .label {{ color: #62748a; font-size: 12px; }}
    .value {{ font-size: 22px; font-weight: 700; margin-top: 4px; }}
    .status {{ color: {status_color}; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th, td {{ padding: 10px 12px; border-bottom: 1px solid #e7edf3; text-align: left; vertical-align: top; }}
    th {{ background: #f8fafc; color: #344054; font-weight: 700; }}
    pre {{ margin: 0; padding: 16px; overflow: auto; max-height: 420px; background: #0f172a; color: #dbeafe; }}
  """


def html_report_build_page(data: dict, packages: list[dict], log_text: str) -> str:
    missing_items = data["missing_items"]
    status = "有缺失" if missing_items else "成功"
    if data["error_lines"]:
        status = "失败"
    if not packages:
        status = "未生成产物"

    missing_rows = [
        [item["project"], item["mode"], item["file_type"], item["reason"]]
        for item in missing_items
    ] or [["-", "-", "-", "没有 MISSING_DELIVER_FILE 记录"]]
    package_rows = [
        [item["name"], html_report_format_size(item["size"])]
        for item in packages
    ] or [["-", "没有在 build/package 下找到产物"]]
    deliver_rows = [
        [item["project"], item["mode"], item["files"]]
        for item in data["deliver_items"]
    ] or [["-", "-", "没有 DELIVER_ITEM 记录"]]
    report_rows = [
        [item["issue"], item["count"]]
        for item in data["report_filled"]
    ] or [["-", "没有 REPORT_FILLED 记录"]]
    error_rows = [[line] for line in data["error_lines"]] or [["没有 Error/Traceback 记录"]]

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <title>Delivery Report</title>
  <style>
    {html_report_style(status)}
  </style>
</head>
<body>
  <header>
    <h1>Delivery Report</h1>
    <div>{html.escape(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))}</div>
  </header>
  <main>
    <div class="summary">
      <div class="card"><div class="label">状态</div><div class="value status">{html.escape(status)}</div></div>
      <div class="card"><div class="label">产物数量</div><div class="value">{len(packages)}</div></div>
      <div class="card"><div class="label">缺失项</div><div class="value">{len(missing_items)}</div></div>
      <div class="card"><div class="label">报告填写</div><div class="value">{len(data["report_filled"])}</div></div>
    </div>
    <section><h2>缺失文件与原因</h2>{html_report_table(["项目", "模块", "缺失项", "原因"], missing_rows)}</section>
    <section><h2>构建产物</h2>{html_report_table(["文件", "大小"], package_rows)}</section>
    <section><h2>交付清单条目</h2>{html_report_table(["项目", "模块", "文件"], deliver_rows)}</section>
    <section><h2>报告填写记录</h2>{html_report_table(["Issue", "填写数量"], report_rows)}</section>
    <section><h2>错误摘要</h2>{html_report_table(["日志"], error_rows)}</section>
    <section><h2>原始日志</h2><pre>{html.escape(log_text)}</pre></section>
  </main>
</body>
</html>
"""


def generate_html_report(log: Path, package_dir: Path, output: Path, error: str = "") -> None:
    log_text = html_report_read_text(log)
    if error:
        log_text += f"\nError: {error}"
    data = html_report_parse_log(log_text)
    packages = html_report_package_files(package_dir)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(html_report_build_page(data, packages, log_text), encoding="utf-8")
    print(f"HTML_REPORT={output}")


# ---------- 命令入口 ----------

def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command")

    html_parser = subparsers.add_parser("html-report")
    html_parser.add_argument("--log", default="build/deliver_output.txt")
    html_parser.add_argument("--package-dir", default="build/package")
    html_parser.add_argument("--output", default="build/html_report/index.html")

    args = parser.parse_args()
    if args.command == "html-report":
        generate_html_report(Path(args.log), Path(args.package_dir), Path(args.output))
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
