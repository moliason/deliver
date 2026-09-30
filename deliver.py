import argparse
import configparser
from urllib.parse import quote
import json
import requests
import subprocess
import os
import re
import glob
import hashlib
import shutil
import sys
import unicodedata
import xml.etree.ElementTree as ET

from fnmatch import fnmatchcase
from io import BytesIO
from redminelib import Redmine
from datetime import datetime, date
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from delivery_report_packager import download_and_fill_feature_reports, has_test_report_deliverable, report_type
from download_file_from_dingtalk import (
    configure_dingtalk, change_application_form, change_application_version_error,
    get_change_application_form, get_peripheral_feature_analysis_form,
    series_requirement_keyword,
)

CONFIG = {}
headers = {}
redmine = None
origin_path = ""
module_list_config_cache = None
repository_match_config_cache = None
change_pinout_contents = {}
architecture_branch_paths = {}
design_reference_paths = {}
XLSX_SUFFIX_GLOB = ".[xX][lL][sS][xX]"
XLS_SUFFIX_GLOB = ".[xX][lL][sS]*"
GPIO_MODE_NAMES = {"gpio", "gpio-afio", "gpio_afio"}
PERIPHERAL_XML_FIELD = "xml"
PERIPHERAL_XML_REPOSITORIES = ("xml", "perihearl_xml", "peripheral_xml", "peripheral_config_information")


def is_excel_filename(filename):
    return fnmatchcase(os.path.splitext(filename)[1], XLS_SUFFIX_GLOB)


def delivery_mode_name(mode):
    return "GPIO" if mode.casefold() in GPIO_MODE_NAMES else mode.upper()


hal_mapping_table = {}
case_mapping_table = {}
example_mapping_table = {}
example_syscfg_mapping_table = {}
LIBRARY_FOLDERS = {}
report_packager_config = {}
issue_info_json = []
pulled = []
folder_AE = {}
folder_SW = {}
TARGET_SERIES_SET = set()


def main():
    global CONFIG, headers, redmine, origin_path, TARGET_SERIES_SET
    global hal_mapping_table, case_mapping_table, example_mapping_table
    global example_syscfg_mapping_table, LIBRARY_FOLDERS, report_packager_config
    parser = argparse.ArgumentParser(description="从 Redmine、GitLab 和钉钉生成交付目录与 Excel 清单")
    parser.add_argument("--config", type=Path, default=Path("delivery.local.json"))
    parser.add_argument("--output", type=Path, default=Path("build/workspace"))
    parser.add_argument("--series", default=os.getenv("TARGET_SERIES", ""), help="只交付指定项目，逗号分隔")
    parser.add_argument("--check-config", action="store_true", help="验证配置，不连接远程服务")
    args = parser.parse_args()
    try:
        config_path = args.config.resolve()
        CONFIG = json.loads(config_path.read_text(encoding="utf-8-sig"))
        gitlab = CONFIG["gitlab"]
        redmine_settings = CONFIG["redmine"]
        gitlab["url"] = os.getenv("GITLAB_URL") or gitlab["url"]
        redmine_settings["url"] = os.getenv("REDMINE_URL") or redmine_settings["url"]
        for name, value in {
            "gitlab.url": gitlab["url"], "gitlab.groups": gitlab["groups"],
            "redmine.url": redmine_settings["url"], "redmine.project_id": redmine_settings["project_id"],
            "GITLAB_TOKEN": os.getenv("GITLAB_TOKEN"), "REDMINE_API_KEY": os.getenv("REDMINE_API_KEY"),
        }.items():
            if not value:
                raise ValueError(f"缺少配置或环境变量: {name}")
        if gitlab.get("clone_protocol", "https") not in {"https", "ssh"}:
            raise ValueError("gitlab.clone_protocol 只能是 https 或 ssh")
        mapping_path = (config_path.parent / CONFIG["selection_mapping"]).resolve()
        if not mapping_path.is_file():
            raise ValueError(f"选型映射文件不存在: {mapping_path}")
        CONFIG["selection_mapping"] = str(mapping_path)
        repositories = CONFIG["repositories"]
        for key in ("series_aliases", "hal", "case", "example", "example_module_paths", "library_roots"):
            repositories[key] = {name.casefold(): value for name, value in repositories[key].items()}
        CONFIG["dingtalk"]["series_aliases"] = {
            name.casefold(): value for name, value in CONFIG["dingtalk"]["series_aliases"].items()
        }
        hal_mapping_table = repositories["hal"]
        case_mapping_table = repositories["case"]
        example_mapping_table = repositories["example"]
        example_syscfg_mapping_table = repositories["example_module_paths"]
        LIBRARY_FOLDERS = {name.upper(): folder for name, folder in repositories["libraries"].items()}
        hal_pattern = re.compile(CONFIG["naming"]["hal_pattern"])
        if hal_pattern.groups != 2:
            raise ValueError("naming.hal_pattern 必须包含模块名、C/H 两个捕获组")
        configure_dingtalk(CONFIG["dingtalk"])
        headers = {"PRIVATE-TOKEN": os.environ["GITLAB_TOKEN"]}
        report_packager_config = {"redmine": redmine_settings}
        TARGET_SERIES_SET = {item.strip() for item in re.split(r"[,，]", args.series) if item.strip()}
        origin_path = str(args.output.resolve())
        if args.check_config:
            print("配置检查通过（未访问远程服务）")
            return 0
        if args.output.exists() and any(args.output.iterdir()):
            raise ValueError("输出目录非空，请用 --output 指定新的交付目录")
        redmine = Redmine(redmine_settings["url"], key=os.environ["REDMINE_API_KEY"], requests={"timeout": 30})
        return deliver()
    except (OSError, ValueError, KeyError) as error:
        print(f"交付失败: {error}", file=sys.stderr)
        return 1

# 从gitlab获取项目URL
def get_git_url_for_repo():
    projects = {}
    gitlab = CONFIG["gitlab"]
    for alias, group in gitlab["groups"].items():
        projects[alias] = []
        page = 1
        while True:
            response = requests.get(
                f"{gitlab['url'].rstrip('/')}/api/v4/groups/{quote(str(group), safe='')}/projects",
                headers=headers,
                params={"per_page": 100, "page": page, "include_subgroups": False},
                timeout=30,
            )
            response.raise_for_status()
            entries = response.json()
            if not isinstance(entries, list):
                raise ValueError(f"GitLab 分组 {alias} 返回格式异常")
            url_field = "ssh_url_to_repo" if gitlab.get("clone_protocol") == "ssh" else "http_url_to_repo"
            projects[alias].extend(item[url_field] for item in entries)
            next_page = response.headers.get("X-Next-Page")
            if not next_page:
                break
            page = int(next_page)
    return projects


#特殊系列处理
def repository_name_keys(name):
    normalized = unicodedata.normalize("NFKC", str(name or "")).casefold().strip()
    if not normalized:
        return set()

    keys = set()

    def add_key(value):
        key = re.sub(r"[^a-z0-9]+", "", value)
        if not key:
            return
        keys.add(key)
        repeated = re.fullmatch(r"(.+?)\1+", key)
        if repeated and repeated.group(1) != key:
            keys.add(repeated.group(1))

    add_key(normalized)
    for bracket_value in re.findall(r"\(([^()]*)\)", normalized):
        add_key(bracket_value)
    return keys


def repository_match_config():
    global repository_match_config_cache
    if repository_match_config_cache is None:
        repository_match_config_cache = {
            "fuzzy_projects": set(),
            "fallbacks": {},
        }
        config = configparser.ConfigParser(interpolation=None)
        config.optionxform = str
        config_path = Path(CONFIG["selection_mapping"])
        try:
            config.read(config_path, encoding="utf-8")
        except (OSError, configparser.Error):
            config = None
        if config:
            if config.has_section("repoMatch"):
                for key, value in config.items("repoMatch"):
                    if key.casefold() == "fuzzyprojects":
                        repository_match_config_cache["fuzzy_projects"].update(
                            item.casefold()
                            for item in re.split(r"[/,，|;]", value)
                            if item.strip()
                        )
            for section in config.sections():
                prefix, separator, project = section.partition(".")
                if section.casefold() == "repofallback":
                    project = "*"
                    prefix = "repoFallback"
                    separator = "."
                if not separator or prefix.casefold() != "repofallback":
                    continue
                target_map = repository_match_config_cache["fallbacks"].setdefault(project.casefold(), {})
                for alias, targets in config.items(section):
                    target_names = [
                        target.strip()
                        for target in re.split(r"[,，|;]", targets)
                        if target.strip()
                    ]
                    for key in repository_name_keys(alias):
                        target_map.setdefault(key, []).extend(target_names)

    return repository_match_config_cache


def repository_fallbacks(project_name, mode):
    config = repository_match_config()
    target_names = []
    for key in repository_name_keys(mode):
        target_names.extend(config["fallbacks"].get("*", {}).get(key, []))
        target_names.extend(config["fallbacks"].get(project_name.casefold(), {}).get(key, []))
    return list(dict.fromkeys(target_names))


def repository_name_from_url(url):
    return re.sub(r"\.git$", "", url.rstrip("/").rsplit("/", 1)[-1], flags=re.IGNORECASE)


def repository_name_matches(requested, actual, allow_fuzzy=False):
    if str(requested).casefold() == str(actual).casefold():
        return True
    if not allow_fuzzy:
        return False

    requested_keys = repository_name_keys(requested)
    actual_keys = repository_name_keys(actual)
    if requested_keys & actual_keys:
        return True

    if any(
        len(requested_key) >= 3 and actual_key.startswith(requested_key)
        for requested_key in requested_keys
        for actual_key in actual_keys
    ):
        return True

    requested_tokens = re.findall(r"[a-z0-9]+", unicodedata.normalize("NFKC", str(requested or "")).casefold())
    actual_tokens = re.findall(r"[a-z0-9]+", unicodedata.normalize("NFKC", str(actual or "")).casefold())
    if len(requested_tokens) == 1 and requested_tokens[0] in actual_tokens:
        return True
    if len(requested_tokens) > 1:
        for index in range(len(actual_tokens) - len(requested_tokens) + 1):
            if actual_tokens[index:index + len(requested_tokens)] == requested_tokens:
                return True
    return False


def find_architecture_repository(project_name, mode, urls):
    exact_matches = [
        url
        for url in urls
        if repository_name_from_url(url).casefold() == mode.casefold()
    ]
    if len(exact_matches) == 1:
        return exact_matches[0]
    if len(exact_matches) > 1:
        print(f"REPO_MATCH_AMBIGUOUS={project_name}\t{mode}\t{','.join(repository_name_from_url(url) for url in exact_matches)}")
        return ""

    fallback_names = repository_fallbacks(project_name, mode)
    if fallback_names:
        fallback_matches = [
            url
            for url in urls
            if any(repository_name_matches(name, repository_name_from_url(url), False) for name in fallback_names)
        ]
        if len(fallback_matches) == 1:
            print(f"REPO_MATCH={project_name}\t{mode}\t{repository_name_from_url(fallback_matches[0])}\tfallback")
            return fallback_matches[0]
        if len(fallback_matches) > 1:
            print(f"REPO_MATCH_AMBIGUOUS={project_name}\t{mode}\t{','.join(repository_name_from_url(url) for url in fallback_matches)}")
        else:
            print(f"REPO_MATCH_MISSING={project_name}\t{mode}")
        return ""

    if project_name.casefold() not in repository_match_config()["fuzzy_projects"]:
        print(f"REPO_MATCH_MISSING={project_name}\t{mode}")
        return ""

    matches = [
        url
        for url in urls
        if repository_name_matches(mode, repository_name_from_url(url), True)
    ]
    if len(matches) == 1:
        matched_name = repository_name_from_url(matches[0])
        print(f"REPO_MATCH={project_name}\t{mode}\t{matched_name}\tfuzzy")
        return matches[0]
    if len(matches) > 1:
        print(f"REPO_MATCH_AMBIGUOUS={project_name}\t{mode}\t{','.join(repository_name_from_url(url) for url in matches)}")
        return ""

    print(f"REPO_MATCH_MISSING={project_name}\t{mode}")
    return ""


# 从gitlab克隆项目
# TODO: 拉取hal库代码
def git_clone(project_name, mode, commit_id, pull_target, projects, check_commit=True, branch="", reference_commit=""):
    os.chdir(origin_path)
    if not os.path.exists(project_name):
        os.mkdir(project_name)
    os.chdir(os.path.join(origin_path, project_name))

    if not os.path.exists(mode):
        os.mkdir(mode)
    os.chdir(os.path.join(origin_path, project_name, mode))

    p_name = ""
    url = ""
    is_peripheral_xml = pull_target.casefold() == PERIPHERAL_XML_FIELD
    mapped_case_repo = case_mapping_table.get(project_name.lower(), "")
    mapped_example_repo = example_mapping_table.get(project_name.lower(), "")
    if "case" in pull_target:
        if not mapped_case_repo:
            print(f"Error: {project_name}未配置Case仓库，跳过case拉取")
            return False
        p_name = CONFIG["repositories"]["hal_group"]
    elif "hal_c" in pull_target or "hal_h" in pull_target or pull_target == "hal_l":
        p_name = CONFIG["repositories"]["hal_group"]
    elif "example" in pull_target and mapped_example_repo:
        p_name = CONFIG["repositories"]["hal_group"]
    elif "ftl" in pull_target or "pin_func_list" in pull_target or "PeriConfig.xml" in pull_target or "tips" in pull_target.lower() or is_peripheral_xml:
        series = CONFIG["repositories"]["series_aliases"].get(project_name.casefold(), project_name.lower())
        p_name = CONFIG["repositories"]["architecture_group"].format(series=series) + "/common"
    else:
        series = CONFIG["repositories"]["series_aliases"].get(project_name.casefold(), project_name.lower())
        p_name = CONFIG["repositories"]["architecture_group"].format(series=series)
    group_urls = next(
        (urls for name, urls in projects.items() if name.casefold() == p_name.casefold()),
        None,
    )
    if group_urls is None:
        print(f"Error: {project_name}找不到仓库分组{p_name}，跳过{pull_target}拉取")
        return False
    if "架构设计" in pull_target:
        url = find_architecture_repository(project_name, mode, group_urls)
    elif is_peripheral_xml:
        for repo_name in PERIPHERAL_XML_REPOSITORIES:
            url = next(
                (item for item in group_urls if repository_name_from_url(item).casefold() == repo_name),
                "",
            )
            if url:
                if repo_name != PERIPHERAL_XML_FIELD:
                    print(f"REPO_MATCH={project_name}\t{PERIPHERAL_XML_FIELD}\t{repo_name}\tfallback")
                break
    else:
        for x in group_urls:
            if "PeriConfig.xml" in pull_target and "peri" in x.split("/")[-1].lower():
                url = x
                break
            elif "ftl" in pull_target and re.search(r"\/ftl\.git", x.lower()):
                url = x
                break
            elif "pin_func_list" in pull_target and re.search(r"\/pin_function_list\.git", x.lower()):
                url = x
                break
            elif "tips" in pull_target.lower() and re.search(r"\/tips\.git", x.lower()):
                url = x
                break
            elif "example" in pull_target and (
                repository_name_from_url(x).casefold() == mapped_example_repo.casefold()
                if mapped_example_repo
                else "example" in repository_name_from_url(x).casefold()
            ):
                url = x
                break
            elif "case" in pull_target and repository_name_from_url(x).casefold() == mapped_case_repo.casefold():
                url = x
                break
            elif pull_target == "hal_l" and repository_name_from_url(x).casefold() == hal_mapping_table.get(project_name.lower(), "").casefold():
                url = x
                break
            elif ("hal_c" in pull_target or "hal_h" in pull_target) and repository_name_from_url(x).casefold() == hal_mapping_table.get(project_name.casefold(), "").casefold():
                url = x
                break

    if url != "":
        repo_name = repository_name_from_url(url)
        if reference_commit:
            reference_commit = str(reference_commit).strip()
            if not re.fullmatch(r"[0-9a-fA-F]{7,40}", reference_commit):
                print(f"ARCH_REFERENCE_PULL_FAILED={project_name}\t{mode}\t{reference_commit}\tcommit id格式不正确")
                return False

            reference_cache_root = os.path.join(origin_path, project_name, mode, "_reference_repos")
            repo_path = os.path.join(reference_cache_root, reference_commit.casefold())
            os.makedirs(reference_cache_root, exist_ok=True)
            try:
                if not os.path.exists(repo_path):
                    subprocess.run(["git", "clone", "--no-checkout", url, repo_path], check=True)
                if not os.path.isdir(os.path.join(repo_path, ".git")):
                    print(f"ARCH_REFERENCE_PULL_FAILED={project_name}\t{mode}\t{reference_commit}\t参考版本缓存目录无效")
                    return False
                subprocess.run(["git", "fetch", "origin"], cwd=repo_path, check=True)
                resolved_commit = subprocess.run(
                    ["git", "rev-parse", "--verify", f"{reference_commit}^{{commit}}"],
                    cwd=repo_path,
                    check=True,
                    capture_output=True,
                    text=True,
                ).stdout.strip()
                subprocess.run(
                    ["git", "checkout", "--detach", "--force", resolved_commit],
                    cwd=repo_path,
                    check=True,
                )
            except subprocess.CalledProcessError:
                print(f"ARCH_REFERENCE_PULL_FAILED={project_name}\t{mode}\t{reference_commit}\tcommit不存在或拉取失败")
                return False
            design_reference_paths[(project_name.casefold(), mode.casefold())] = repo_path
            os.chdir(repo_path)
        elif branch:
            branch_check = subprocess.run(
                ["git", "check-ref-format", "--branch", branch],
                capture_output=True,
                text=True,
            )
            if branch_check.returncode != 0:
                print(f"ARCH_BRANCH_PULL_FAILED={project_name}\t{mode}\t{branch}\t分支名不合法")
                return False

            branch_folder = branch.replace("/", "__").replace("\\", "__")
            branch_cache_root = os.path.join(origin_path, project_name, mode, "_branch_repos")
            repo_path = os.path.join(branch_cache_root, branch_folder)
            os.makedirs(branch_cache_root, exist_ok=True)
            if not os.path.exists(repo_path):
                try:
                    subprocess.run(
                        ["git", "clone", "--branch", branch, "--single-branch", url, repo_path],
                        check=True,
                    )
                except subprocess.CalledProcessError:
                    print(f"ARCH_BRANCH_PULL_FAILED={project_name}\t{mode}\t{branch}\t分支不存在或克隆失败")
                    return False
            if not os.path.isdir(repo_path):
                print(f"ARCH_BRANCH_PULL_FAILED={project_name}\t{mode}\t{branch}\t分支缓存目录无效")
                return False
            current_branch = subprocess.run(
                ["git", "branch", "--show-current"],
                cwd=repo_path,
                capture_output=True,
                text=True,
            )
            if current_branch.returncode != 0 or current_branch.stdout.strip() != branch:
                print(f"ARCH_BRANCH_PULL_FAILED={project_name}\t{mode}\t{branch}\t分支缓存与目标分支不一致")
                return False
            architecture_branch_paths[(project_name.casefold(), mode.casefold(), branch)] = repo_path
            os.chdir(repo_path)
        else:
            repo_paths = [
                os.path.join(origin_path, project_name, mode, repo_name),
                os.path.join(origin_path, project_name, mode, repo_name.lower()),
                os.path.join(origin_path, project_name, mode, repo_name.upper()),
            ]
            if "架构设计" in pull_target and repo_name.casefold() == "usb":
                repo_paths.extend(glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project_name), "*", repo_name)))
                repo_paths.extend(glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project_name), "*", repo_name.lower())))
                repo_paths.extend(glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project_name), "*", repo_name.upper())))
            if "ftl" in pull_target or "pin_func_list" in pull_target or "tips" in pull_target.lower():
                repo_paths.extend(glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project_name), "*", repo_name)))
                repo_paths.extend(glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project_name), "*", repo_name.lower())))
                repo_paths.extend(glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project_name), "*", repo_name.upper())))

            if not any(os.path.exists(repo_path) for repo_path in repo_paths):
                subprocess.run(['git', 'clone', url], check=True)

            for repo_path in repo_paths:
                if os.path.exists(repo_path):
                    os.chdir(repo_path)
                    break
            else:
                raise FileNotFoundError(os.path.join(origin_path, project_name, mode, repo_name))
        if url not in pulled:
            pulled.append(url)

        if "架构设计" in pull_target:
            # 先修改拉取目录中的界面表/xml文件名，后续交付清单和打包会自动使用带版本号的新名称。
            rename_interface_tables(repo_path)
            rename_config_xml_files(repo_path)
        elif "tips" in pull_target.lower():
            rename_interface_tables(repo_path)

        if branch:
            head_commit = subprocess.run(
                ["git", "rev-parse", "HEAD"],
                cwd=repo_path,
                check=True,
                capture_output=True,
                text=True,
            ).stdout.strip()
            print(f"ARCH_BRANCH_PULL={project_name}\t{mode}\t{branch}\t{head_commit}")
        elif reference_commit:
            print(f"ARCH_REFERENCE_PULL={project_name}\t{mode}\t{resolved_commit}")

        # 普通交付需要校验 Redmine 填写的 commit id；Change 包不信任该字段，
        # 后续会直接用仓库里的 HEAD 和 HEAD^ 计算引脚差异，所以这里允许跳过校验。
        if check_commit and not branch and not reference_commit and "example" not in pull_target and "tips" not in pull_target.lower():
            if commit_id != "" and commit_id != "000000":
                result = subprocess.run(
                ['git', 'rev-parse', 'HEAD'],
                cwd=os.getcwd(),
                check=True,
                capture_output=True,
                text=True
                )
                if result.stdout.strip() != commit_id:
                    return False
            else:
                return False

        return True
    else:
        return False


def architecture_remote_branches(repo_path):
    try:
        branches_output = subprocess.run(
            ["git", "ls-remote", "--symref", "origin", "HEAD", "refs/heads/*"],
            cwd=repo_path,
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        ).stdout
    except (OSError, subprocess.CalledProcessError) as error:
        print(f"ARCH_BRANCH_LIST_FAILED={repo_path}\t{error}")
        return "", None

    default_branch = ""
    branches = set()
    for line in branches_output.splitlines():
        parts = line.split()
        if len(parts) >= 3 and parts[0] == "ref:" and parts[2] == "HEAD":
            default_branch = parts[1].removeprefix("refs/heads/")
        elif len(parts) >= 2 and parts[1].startswith("refs/heads/"):
            branches.add(parts[1].removeprefix("refs/heads/"))
    if not default_branch:
        print(f"ARCH_BRANCH_LIST_FAILED={repo_path}\t远程仓库未设置默认分支")
        return "", None

    branches = sorted(branches, key=str.casefold)
    print(f"ARCH_BRANCH_LIST={repo_path}\tdefault={default_branch}\tbranches={','.join(branches)}")
    return default_branch, branches

# 从redmine的issue中下载附件
def download_url(project_name, mode, attachment, filename_filter=None):
    os.chdir(origin_path)
    if not os.path.exists(project_name):
        os.mkdir(project_name)
    os.chdir(os.path.join(origin_path, project_name))

    if not os.path.exists(mode):
        os.mkdir(mode)
    os.chdir(os.path.join(origin_path, project_name, mode))

    for key, value in attachment.items():
        if filename_filter and not filename_filter(key):
            continue
        response = requests.get(value, stream=True)
        response.raise_for_status()  # 检查请求是否成功

        with open("./"+key, 'wb') as file:
            for chunk in response.iter_content(chunk_size=8192):
                file.write(chunk)

def delete_empty_folder(src):
    """
    递归删除空文件夹
    """
    if not os.path.exists(src):
        return

    # 使用 bottom-up 方式遍历，从最深的目录开始
    for root, dirs, files in os.walk(src, topdown=False):
        # 先处理子目录
        for dir_name in dirs:
            dir_path = os.path.join(root, dir_name)
            try:
                if not os.listdir(dir_path):  # 检查目录是否为空
                    os.rmdir(dir_path)
                    print(f"已删除空目录: {dir_path}")
            except (OSError, FileNotFoundError):
                # 如果目录不存在或无法删除（可能已被删除），则跳过
                continue

        # 再检查当前根目录是否为空
        try:
            if not os.listdir(root):
                os.rmdir(root)
                print(f"已删除空目录: {root}")
        except (OSError, FileNotFoundError):
            # 如果根目录不存在或无法删除（可能已被删除），则跳过
            continue

def copy_files(src, dst):
    for root, dirs, files in os.walk(src):
        dirs[:] = [d for d in dirs if d != ".git"]

        for file in files:
            shutil.copy2(os.path.join(root, file), dst)


def copy_directory_contents(src, dst):
    # 架构设计交付必须保留 Git 仓库内的目录层级，例如 ADC0、ADC1 不能被拍平。
    if not os.path.isdir(src):
        return
    shutil.copytree(
        src,
        dst,
        dirs_exist_ok=True,
        ignore=shutil.ignore_patterns(".git"),
    )
    # 仓库根目录的公共 .c/.h 是公共代码，保持在外层即可，不再复制到每个子模块目录。


def module_dir_names(path):
    # 仅用于识别真正的子模块目录；黑名单目录仍会被复制，只是不作为 Peripheral 行。
    ignored_dirs = {".git", "备注"} | module_dir_blacklist()
    return [
        name
        for name in os.listdir(path)
        if name.casefold() not in ignored_dirs and os.path.isdir(os.path.join(path, name))
    ]


def design_repo_path(project, mode):
    expected_path = os.path.join(origin_path, project, mode, mode)
    mode_path = os.path.join(origin_path, project, mode)
    fallback_names = repository_fallbacks(project, mode)
    expected_names = [mode] + fallback_names
    fuzzy_enabled = project.casefold() in repository_match_config()["fuzzy_projects"]
    if os.path.isdir(mode_path):
        for expected_name in expected_names:
            expected_keys = repository_name_keys(expected_name)
            matches = sorted(
                os.path.join(mode_path, name)
                for name in os.listdir(mode_path)
                if os.path.isdir(os.path.join(mode_path, name))
                and repository_name_keys(name) & expected_keys
            )
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                print(f"REPO_PATH_AMBIGUOUS={project}\t{mode}\t{','.join(matches)}")
                return expected_path

        if fuzzy_enabled:
            matches = sorted(
                os.path.join(mode_path, name)
                for name in os.listdir(mode_path)
                if os.path.isdir(os.path.join(mode_path, name))
                and repository_name_matches(mode, name, True)
            )
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                print(f"REPO_PATH_AMBIGUOUS={project}\t{mode}\t{','.join(matches)}")
                return expected_path

    if fallback_names:
        project_path = os.path.join(origin_path, project)
        if os.path.isdir(project_path):
            for parent_name in os.listdir(project_path):
                parent_path = os.path.join(project_path, parent_name)
                if not os.path.isdir(parent_path):
                    continue
                if not any(repository_name_keys(parent_name) & repository_name_keys(name) for name in fallback_names):
                    continue
                for child_name in os.listdir(parent_path):
                    candidate = os.path.join(parent_path, child_name)
                    if os.path.isdir(candidate) and repository_name_matches(mode, child_name, fuzzy_enabled):
                        return candidate
    return expected_path


def hal_deliverable_info(deliverable):
    match = re.fullmatch(CONFIG["naming"]["hal_pattern"], deliverable.strip(), re.IGNORECASE)
    if not match:
        return None
    return match.group(1).lower(), match.group(2).upper()


def copy_hal_files(project, mode, deliverable, hal_info, mode_package_path, missing_items):
    peripheral, file_kind = hal_info
    repo_name = hal_mapping_table.get(project.lower(), "")
    if not repo_name:
        record_missing(missing_items, project, mode, deliverable)
        return

    source_dir = "Include" if file_kind == "H" else "Source"
    extension = "h" if file_kind == "H" else "c"
    target_dir = os.path.join(mode_package_path, f"{delivery_mode_name(mode)}_hal_peripheral", source_dir)
    hal_files = glob.glob(os.path.join(
        origin_path, project, mode, repo_name, "Firmware", "*_hal_peripheral",
        source_dir, f"*{peripheral}*.{extension}",
    ))
    os.makedirs(target_dir, exist_ok=True)
    for file in hal_files:
        shutil.copy2(file, target_dir)
    if not glob.glob(os.path.join(glob.escape(target_dir), f"*{peripheral}*.{extension}")):
        record_missing(missing_items, project, mode, deliverable)


def copy_hal_library(project, mode, deliverable, mode_package_path, missing_items):
    format_reason = deliverable_format_reason(deliverable, mode, project)
    if format_reason:
        record_missing(missing_items, project, mode, f"交付物清单格式:{deliverable}", format_reason)
        return

    folder_name = LIBRARY_FOLDERS[deliverable.strip().upper()]
    source_path = os.path.join(
        origin_path, project, mode, hal_mapping_table[project.lower()],
        *Path(CONFIG["repositories"]["library_roots"][project.casefold()]).parts, folder_name,
    )
    if not has_any_file(source_path):
        record_missing(missing_items, project, mode, deliverable, f"HAL仓库中未找到 {folder_name} 或目录为空")
        return
    copy_directory_contents(source_path, os.path.join(mode_package_path, "middleware", folder_name))


def copy_peripheral_xml_files(project, mode, target_path, missing_items):
    source_parent = os.path.join(origin_path, project, mode)
    repo_dirs = {}
    if os.path.isdir(source_parent):
        repo_dirs = {
            name.casefold(): os.path.join(source_parent, name)
            for name in os.listdir(source_parent)
            if os.path.isdir(os.path.join(source_parent, name))
        }
    source_path = next(
        (repo_dirs[name] for name in PERIPHERAL_XML_REPOSITORIES if name in repo_dirs),
        "",
    )
    if source_path:
        copy_directory_contents(source_path, target_path)
    if not any(
        name.lower().endswith(".xml")
        for _, _, files in os.walk(target_path)
        for name in files
    ):
        record_missing(missing_items, project, mode, PERIPHERAL_XML_FIELD)


def copy_architecture_branch(project, mode, branch, mode_package_path, missing_items):
    source_path = architecture_branch_paths.get((project.casefold(), mode.casefold(), branch), "")
    if not os.path.isdir(source_path):
        record_missing(
            missing_items,
            project,
            mode,
            f"架构设计分支:{branch}",
            f"架构设计仓库分支 {branch} 不存在",
        )
        return

    branch_folder = branch.replace("/", "__").replace("\\", "__")
    target_path = os.path.join(mode_package_path, "branches", branch_folder)
    copy_directory_contents(source_path, target_path)
    if mode.casefold() in GPIO_MODE_NAMES:
        for root, _, files in os.walk(target_path):
            for file in files:
                if "引脚" in file and file.lower().endswith(".xlsx"):
                    file_path = os.path.join(root, file)
                    os.remove(file_path)
                    print(f"GPIO_PIN_XLSX_REMOVED={file_path}")
    check_middleware_filenames(target_path, project, mode, missing_items)
    if missing_design_file_types(target_path):
        record_missing(
            missing_items,
            project,
            mode,
            f"架构设计分支:{branch}",
            f"架构设计仓库分支 {branch} 的打包目录为空",
        )
        return
    check_design_md5_files(target_path, project, mode, missing_items)
    print(f"ARCH_BRANCH_PACKAGE={project}\t{mode}\t{branch}\t{target_path}")


def missing_design_file_types(path):
    if os.path.isdir(path):
        ignored_dirs = {".git"} | module_dir_blacklist()
        for root, dirs, files in os.walk(path):
            if {part.casefold() for part in Path(root).parts} & ignored_dirs:
                dirs[:] = []
                continue
            dirs[:] = [directory for directory in dirs if directory.casefold() not in ignored_dirs]
            for file in files:
                if not file.startswith("."):
                    return []
    return ["file"]

#检查中间件文件命名规范
def check_middleware_filenames(path, project, mode, missing_items):
    valid_name_parts = (
        "_中间件embedded builder界面表_",
        "_中间件工程配置excel_",
    )
    for root, dirs, files in os.walk(path):
        dirs[:] = [directory for directory in dirs if directory != ".git"]
        for filename in files:
            normalized_name = filename.casefold()
            if "中间件" not in filename or not normalized_name.endswith(".xlsx"):
                continue
            if any(name_part in normalized_name for name_part in valid_name_parts):
                continue
            record_missing(
                missing_items,
                project,
                mode,
                f"中间件命名:{filename}",
                "中间件文件命名不规范，应包含 _中间件Embedded Builder界面表_ 或 _中间件工程配置Excel_",
            )


def check_design_md5_files(path, project, mode, missing_items):
    ignored_dirs = {".git"} | module_dir_blacklist()
    # 选型配置与脚本放在一起，Jenkins 不依赖本机的 XBuilder.ini。
    selection_config_path = Path(CONFIG["selection_mapping"])
    selection_config = configparser.ConfigParser(interpolation=None)
    selection_config.optionxform = str
    selection_config_error = ""
    try:
        with open(selection_config_path, "r", encoding="utf-8-sig") as file:
            selection_config.read_file(file)
    except (OSError, configparser.Error) as error:
        selection_config_error = str(error)

    configured_series = {}
    series_parts = {}
    series_patterns = {}
    project_series = []
    if not selection_config_error and not selection_config.has_section("seriesParts"):
        selection_config_error = "缺少 [seriesParts] 配置段"
    if not selection_config_error:
        for configured_name, configured_parts in selection_config.items("seriesParts"):
            series = configured_name.strip()
            series_key = series.casefold()
            configured_series[series_key] = series
            series_parts[series_key] = [
                item.strip()
                for item in re.split(r"[/,，]", configured_parts)
                if item.strip()
            ]

            series_name = re.sub(r"_std$", "", series, flags=re.IGNORECASE)
            family_match = re.match(rf"^({re.escape(CONFIG['naming']['series_prefix'])}[A-Za-z]+)", series_name, re.IGNORECASE)
            family_prefix = family_match.group(1) if family_match else CONFIG["naming"]["series_prefix"]
            series_patterns[series_key] = []
            for option in series_name.split("_"):
                if not option.casefold().startswith(CONFIG["naming"]["series_prefix"].casefold()):
                    option = (CONFIG["naming"]["series_prefix"] if option[0].isalpha() else family_prefix) + option
                series_patterns[series_key].append(option.casefold().replace("x", "?") + "*")

        if not configured_series:
            selection_config_error = "[seriesParts] 中未配置 MCU 系列"

        project_aliases = {}
        if selection_config.has_section("projectAlias"):
            project_aliases = {
                name.casefold(): value
                for name, value in selection_config.items("projectAlias")
            }
        project_series = [
            item.strip()
            for item in re.split(r"[/,，]", project_aliases.get(project.casefold(), ""))
            if item.strip()
        ]
        if not project_series and project.casefold() in configured_series:
            project_series = [configured_series[project.casefold()]]
    config_error_reported = False

    for root, dirs, files in os.walk(path):
        root_dirs = {part.casefold() for part in Path(root).parts}
        if root_dirs & ignored_dirs:
            dirs[:] = []
            continue
        dirs[:] = [directory for directory in dirs if directory.casefold() not in ignored_dirs]

        md5_files = [
            os.path.join(root, name)
            for name in files
            if "md5" in name.lower() and name.lower().endswith(".txt")
        ]
        if not md5_files:
            continue

        # MD5 只校验可控范围：MD5 同目录文件，以及模块最外层的公共 .c/.h。
        path_tokens = [
            match.group(0).casefold()
            for directory in Path(os.path.relpath(root, path)).parts
            for match in re.finditer(rf"{re.escape(CONFIG['naming']['series_prefix'])}[A-Za-z0-9]+", directory, re.IGNORECASE)
        ]
        path_series = [
            series
            for series_key, series in configured_series.items()
            if any(
                part.casefold().startswith(token)
                for part in series_parts.get(series_key, [])
                for token in path_tokens
            )
        ]
        expected_series = path_series or project_series
        expected_series_keys = {series.casefold() for series in expected_series}
        md5_lines = []
        for md5_file in md5_files:
            current_md5_lines = []
            for encoding in ["utf-8-sig", "gbk"]:
                try:
                    with open(md5_file, "r", encoding=encoding) as file:
                        current_md5_lines = file.readlines()
                    break
                except UnicodeDecodeError:
                    continue
            md5_lines.extend(current_md5_lines)

            # MD5 文件末尾记录了本次界面表生成时的 MCU 选型；从后往前取最后一条，避免旧记录干扰。
            md5_fields = {}
            for line in reversed(current_md5_lines):
                match = re.match(r"^\s*(MCUSeries|MCUPart)\s*[:=]\s*(.*?)\s*$", line, re.IGNORECASE)
                if match:
                    md5_fields.setdefault(match.group(1).casefold(), match.group(2).strip())
                if len(md5_fields) == 2:
                    break

            md5_name = os.path.basename(md5_file)
            md5_series = md5_fields.get("mcuseries", "")
            md5_part = md5_fields.get("mcupart", "")
            if selection_config_error:
                if not config_error_reported:
                    record_missing(
                        missing_items,
                        project,
                        mode,
                        "MD5选型配置",
                        f"无法读取 {selection_config_path}: {selection_config_error}",
                    )
                    config_error_reported = True
                continue

            selection_result = "通过"
            selection_reason = ""
            md5_series_key = md5_series.casefold()
            if not md5_series:
                # 未提供选型字段时只保留检查日志，不加入最终 missing_file。
                selection_result = "跳过（未找到 MCUSeries）"
            elif md5_series_key not in configured_series:
                selection_result = "MCUSeries 未配置"
                selection_reason = f"MD5 文件 MCUSeries={md5_series} 不在配置文件的 MCUSeries 列表中"
            elif expected_series_keys and md5_series_key not in expected_series_keys:
                selection_result = "MCUSeries 不匹配"
                selection_reason = (
                    f"MD5 文件 MCUSeries={md5_series}，交付工程/子系列目录应为 {'/'.join(expected_series)}"
                )
            else:
                allowed_parts = series_parts.get(md5_series_key, [])
                directory_parts = [
                    part
                    for part in allowed_parts
                    if any(part.casefold().startswith(token) for token in path_tokens)
                ]
                allowed_parts = directory_parts or allowed_parts
                md5_part_key = md5_part.casefold()
                part_matches = md5_part_key in {part.casefold() for part in allowed_parts}
                if md5_part and not part_matches:
                    if directory_parts:
                        part_matches = any(md5_part_key.startswith(token) for token in path_tokens)
                    else:
                        # 名单右侧是代表型号，同时支持系列名中的 x 通配。
                        part_matches = any(
                            fnmatchcase(md5_part_key, pattern)
                            for pattern in series_patterns.get(md5_series_key, [])
                        )
                if allowed_parts and not md5_part:
                    selection_result = "缺少 MCUPart"
                    selection_reason = f"MCUSeries={md5_series} 已配置具体选型，但 MD5 文件未找到 MCUPart"
                elif allowed_parts and not part_matches:
                    selection_result = "MCUPart 不匹配"
                    selection_reason = f"MD5 文件 MCUPart={md5_part} 不属于 MCUSeries={md5_series}"

            if selection_reason:
                record_missing(missing_items, project, mode, f"MD5选型:{md5_name}", selection_reason)
            if not expected_series:
                selection_result += "（未配置交付工程别名）"
            print(
                f"MD5_SELECTION_CHECK={project}\t{mode}\t{md5_name}\t"
                f"MCUSeries={md5_series or '<空>'}\tMCUPart={md5_part or '<空>'}\t{selection_result}"
            )

        md5_values = {md5.lower() for line in md5_lines for md5 in re.findall(r"\b[a-fA-F0-9]{32}\b", line)}

        target_files = {
            name: os.path.join(root, name)
            for name in files
            if name.lower().endswith((".c", ".h", ".xml"))
        }
        if os.path.abspath(root) != os.path.abspath(path):
            for name in os.listdir(path):
                current_path = os.path.join(path, name)
                if os.path.isfile(current_path) and name.lower().endswith((".c", ".h")):
                    target_files.setdefault(name, current_path)
            for name in os.listdir(root):
                current_path = os.path.join(root, name)
                if os.path.isdir(current_path):
                    for child_root, _, child_files in os.walk(current_path):
                        for child_name in child_files:
                            if child_name.lower().endswith((".c", ".h", ".xml")):
                                target_files.setdefault(child_name, os.path.join(child_root, child_name))

        for name, current_path in target_files.items():
            with open(current_path, "rb") as file:
                data = file.read()
            data = data.replace(b"\r\n", b"\n").replace(b"\r", b"\n").replace(b"\n", b"\r\n")
            actual_md5 = hashlib.md5(data).hexdigest()
            if actual_md5 in md5_values:
                continue
            record_missing(missing_items, project, mode, f"MD5校验:{name}", f"{name} 的实际 MD5 {actual_md5} 不在 MD5 文件记录中")


def missing_reason(file_type):
    if file_type.startswith("架构设计:"):
        missing_type = file_type.split(":", 1)[1]
        if missing_type == "file":
            return "交付物清单要求架构设计，但打包目录为空"
        return f"架构设计仓库已拉取，但缺少必需文件类型：{missing_type}"

    reasons = {
        "example": "交付物清单要求 Example，但 example 仓库目录为空或目录名未匹配到模块",
        "测试报告": "交付物清单要求测试报告，但 Redmine 附件或下载目录中未找到测试报告 xls/xlsx/xlsm",
        "评审表": "交付物清单要求评审表，但 Redmine 附件或下载目录中未找到评审表 xlsx",
        "不合格反馈单": "交付物清单要求反馈单，但 Redmine 附件或下载目录中未找到不合格反馈单 xlsx",
        "变更申请单": "交付物清单要求变更申请单，但钉盘/下载目录中未找到变更申请单 xlsx",
        "外设功能分析表": "交付物清单要求外设功能分析表，但钉盘/下载目录中未找到外设功能分析表 xlsx",
        "引脚功能缺失清单": "交付物清单要求引脚功能缺失清单，但钉盘/下载目录中未找到引脚功能缺失清单 xlsx",
        "未实现备注": "交付物清单要求未实现备注，但钉盘/下载目录中未找到对应 xlsx",
        "QA": "交付物清单要求 QA 文件，但钉盘/下载目录中未找到对应 xlsx",
        "Xbuilder": "交付物清单要求 Xbuilder 文件，但钉盘/下载目录中未找到对应 xlsx",
        "Firmware": "交付清单要求 Firmware，但未找到对应 HAL 仓库或仓库为空",
        "main.ftl": "交付物清单要求 main.ftl，但 main&libopt 仓库未拉到或仓库内缺少 main.ftl",
        "libopt.ftl": "交付物清单要求 libopt.ftl，但 main&libopt 仓库未拉到或仓库内缺少 libopt.ftl",
        "pin_function_list": "交付物清单要求 pin_function_list，但 pin 仓库未拉到或仓库内缺少对应 xlsx",
    }
    return reasons.get(file_type, "交付物清单要求该文件，但对应仓库/目录中未找到匹配文件")


#缺失文件记录与打印
def record_missing(missing_items, project, mode, file_type, reason=""):
    missing_items.append({
        "project": project,
        "mode": mode,
        "file_type": file_type,
        "reason": reason or missing_reason(file_type),
    })


def print_missing_items(missing_items):
    for item in missing_items:
        print(f"MISSING_DELIVER_FILE={item['project']}\t{item['mode']}\t{item['file_type']}\t{item['reason']}")


def architecture_branch_request(deliverable):
    item = deliverable.strip()
    if item == "架构设计":
        return "default", []
    if item.casefold() == "架构设计@all":
        return "all", []
    if item.startswith("架构设计@"):
        return "default", [branch.strip() for branch in item.split("@")[1:]]
    match = re.fullmatch(r"(?:架构设计（(.*)）|架构设计\((.*)\))", item)
    if match:
        branch = match.group(1) if match.group(1) is not None else match.group(2)
        return "only", [branch.strip()]
    return None


def deliverable_format_reason(deliverable, mode, project):
    item = deliverable.strip()
    code_series = re.sub(rf"^{re.escape(CONFIG['naming']['series_prefix'])}", "", project, flags=re.IGNORECASE)
    has_test_keyword = "系统测试" in item or "集成测试" in item
    if has_test_keyword and "testreport" not in item.lower() and "测试报告" not in item:
        return "测试报告交付物格式不规范，应写成 系统测试TestReport/系统测试报告，不能只写 系统测试"

    architecture_request = architecture_branch_request(item)
    if architecture_request is not None:
        request_mode, branches = architecture_request
        if request_mode == "all":
            return ""
        if any(not branch for branch in branches):
            return "架构设计分支格式不规范，应写成 架构设计@分支名 或 架构设计（分支名）"
        if len(branches) > 1 and any(branch.casefold() == "all" for branch in branches):
            return "架构设计@all 不能与其他分支名同时填写"
        for branch in branches:
            if subprocess.run(
                ["git", "check-ref-format", "--branch", branch],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            ).returncode != 0:
                return f"架构设计分支名不合法：{branch}"
        return ""
    if "架构设计" in item:
        return "架构设计交付物格式不规范，应写成 架构设计、架构设计@分支名 或 架构设计（分支名）"

    if item.upper() in LIBRARY_FOLDERS:
        if project.casefold() not in CONFIG["repositories"]["library_roots"]:
            return "该项目未配置中间件 library_roots"
        if project.casefold() not in hal_mapping_table:
            return "该项目未配置 HAL 仓库"
        return ""

    # 代码类交付物只按现有拉取规则检查模块是否对应，避免把合法的 CAN/DMA 写法误报。
    hal_info = hal_deliverable_info(item)
    if hal_info:
        if hal_info[0].upper() != mode.upper():
            return f"HAL代码交付物模块不匹配，当前模块为 {mode}，请按 naming.hal_pattern 填写"
        return ""

    example_match = re.fullmatch(r"([A-Za-z0-9]+)_E([A-Z0-9]+)", item, re.IGNORECASE)
    if example_match:
        if example_match.group(1).casefold() != code_series.casefold() or example_match.group(2).upper() != mode.upper():
            return f"Example交付物系列或模块不匹配，{project} 的 {mode} 模块应填写 {code_series}_E{mode}"
        return ""

    case_match = re.fullmatch(r"([A-Za-z0-9]+)_C([A-Z0-9]+)", item, re.IGNORECASE)
    if case_match:
        if case_match.group(1).casefold() != code_series.casefold() or case_match.group(2).upper() != mode.upper():
            return f"Case代码交付物系列或模块不匹配，{project} 的 {mode} 模块应填写 {code_series}_C{mode}"
        return ""

    if item.upper() != "EXAMPLE" and (
        re.search(r"(?:^|_)example(?:_|$)", item, re.IGNORECASE)
        or re.fullmatch(r"[A-Za-z0-9]+_E[A-Z0-9]+_[CH]", item, re.IGNORECASE)
    ):
        return f"Example交付物格式不规范，{project} 的 {mode} 模块应填写 {code_series}_E{mode}，不需要 C/H 后缀"

    wrong_code_match = re.fullmatch(r"[A-Za-z0-9]+_([A-Z0-9]+)_[ECH]", item, re.IGNORECASE)
    if wrong_code_match:
        return f"代码类交付物格式不规范，例如 Example 填写 {code_series}_E{mode}，Case 填写 {code_series}_C{mode}"
    return ""


def has_any_file(path):
    if not os.path.isdir(path):
        return False

    for _, _, files in os.walk(path):
        if files:
            return True
    return False


def packaged_examples(deliver_items:dict, missing_items:list):
    for project in deliver_items:
        code_series = re.sub(rf"^{re.escape(CONFIG['naming']['series_prefix'])}", "", project, flags=re.IGNORECASE)
        mapped_example_repo = example_mapping_table.get(project.lower(), "")
        os.chdir(origin_path)
        root_path = f"[{project}]Example {datetime.now().strftime('%Y%m%d')}"
        package_path = os.path.join(origin_path, root_path)
        os.mkdir(root_path)
        os.chdir(root_path)
        rows = []
        full_example_rows_written = False

        for mode in sorted(deliver_items[project]):
            mode_deliverables = deliver_items[project][mode]
            mode_name = mode.rsplit("_", 1)[0] if mode.endswith(("_SW", "_AE")) else mode
            has_full_example = any(item.strip().upper() == "EXAMPLE" for item in mode_deliverables)
            has_module_example = any(re.fullmatch(rf"{re.escape(code_series)}_E[A-Z0-9]+", item.strip(), re.IGNORECASE) for item in mode_deliverables)
            has_example = has_full_example or has_module_example
            if not has_example:
                continue
            merge_into_ae = mode.endswith("_AE") and has_module_example and not has_full_example
            if merge_into_ae:
                mode_path = os.path.join(
                    origin_path,
                    f"[{project}]Deliver List To AE {datetime.now().strftime('%Y%m%d')}",
                    delivery_mode_name(mode_name),
                    f"{delivery_mode_name(mode_name)}_Example",
                )
                os.makedirs(mode_path, exist_ok=True)
            else:
                mode_path = package_path if has_full_example and not has_module_example else os.path.join(package_path, delivery_mode_name(mode_name))
                if mode_path != package_path:
                    os.mkdir(mode_path)

            source_path = os.path.join(origin_path, project, mode_name)
            examples = [
                os.path.join(source_path, name)
                for name in os.listdir(source_path)
                if "example" in name.lower() or name.casefold() == mapped_example_repo.casefold()
            ]
            if has_module_example and not has_full_example:
                examples = []
                fallback_path = example_syscfg_mapping_table.get(mode_name.casefold())
                for name in os.listdir(source_path):
                    if "example" not in name.lower() and name.casefold() != mapped_example_repo.casefold():
                        continue
                    example_path = os.path.join(source_path, name, mode_name.upper())
                    if not os.path.isdir(example_path) and fallback_path:
                        example_path = os.path.join(source_path, name, *fallback_path)
                    if os.path.isdir(example_path):
                        examples.append(example_path)
            for example in examples:
                copy_directory_contents(example, mode_path)
            if not has_any_file(mode_path):
                record_missing(missing_items, project, mode_name, "example")

            if merge_into_ae:
                continue

            if has_full_example and not has_module_example:
                if not full_example_rows_written:
                    ignored_dirs = {".git", "备注"} | module_dir_blacklist()
                    for module in sorted(
                        name
                        for name in os.listdir(package_path)
                        if name.casefold() not in ignored_dirs and os.path.isdir(os.path.join(package_path, name))
                    ):
                        rows.append({
                            "Peripheral": module.upper(),
                            "IDE界面表模板": "",
                            "标准驱动": "NO",
                            "xx_config.xml": "-",
                            "MD5": "NO",
                            "其他": "Example",
                        })
                        print(f"DELIVER_ITEM={project}\t{module.upper()}\tExample")
                    full_example_rows_written = True
                continue

            if has_any_file(mode_path):
                rows.append({
                    "Peripheral": delivery_mode_name(mode_name),
                    "IDE界面表模板": "",
                    "标准驱动": "NO",
                    "xx_config.xml": "-",
                    "MD5": "NO",
                    "其他": "Example",
                })
                print(f"DELIVER_ITEM={project}\t{mode_name}\tExample")
        delete_empty_folder(package_path)

        if os.path.isdir(package_path):
            list_path = next_version_path(f"{package_path}.xlsx")
            write_deliver_list_workbook(rows, list_path)
            print(f"DELIVER_LIST={list_path}")


def packaged_firmwares(firmware_projects:set, missing_items:list):
    for project in sorted(firmware_projects):
        os.chdir(origin_path)
        root_path = f"[{project}]Firmware {datetime.now().strftime('%Y%m%d')}"
        package_path = os.path.join(origin_path, root_path)
        os.mkdir(root_path)
        rows = []

        source_path = os.path.join(origin_path, project, "Firmware")
        if os.path.isdir(source_path):
            for repo_name in sorted(os.listdir(source_path)):
                repo_path = os.path.join(source_path, repo_name)
                if not os.path.isdir(repo_path):
                    continue
                shutil.copytree(
                    repo_path,
                    os.path.join(package_path, repo_name),
                    ignore=shutil.ignore_patterns(".git"),
                )
                rows.append({
                    "Peripheral": repo_name,
                    "IDE界面表模板": "",
                    "标准驱动": "NO",
                    "xx_config.xml": "-",
                    "MD5": "NO",
                    "其他": "Firmware",
                })
                print(f"DELIVER_ITEM={project}\t{repo_name}\tFirmware")

        delete_empty_folder(package_path)
        if not os.path.isdir(package_path):
            record_missing(missing_items, project, "Firmware", "Firmware")
            continue

        list_path = next_version_path(f"{package_path}.xlsx")
        write_deliver_list_workbook(rows, list_path)
        print(f"DELIVER_LIST={list_path}")

# 打包交付物_SW
def packaged_deliverables_SW(deliver_items:dict, missing_items:list, package_title="Deliver List"):
    package_paths = []
    for project in deliver_items:
        code_series = re.sub(rf"^{re.escape(CONFIG['naming']['series_prefix'])}", "", project, flags=re.IGNORECASE)
        sw_modes = [
            mode
            for mode in deliver_items[project]
            if mode.endswith("_SW")
        ]
        if not sw_modes:
            continue

        os.chdir(origin_path)
        root_path = f"[{project}]{package_title} To SW {datetime.now().strftime('%Y%m%d')}"
        os.mkdir(root_path)
        os.chdir(root_path)
        os.mkdir("Common")
        common_path = os.path.join(origin_path, root_path, "Common")

        for mode in sw_modes:
            mode_name = mode.rsplit("_", 1)[0]
            mode_package_path = os.path.join(origin_path, root_path, delivery_mode_name(mode_name))
            os.makedirs(mode_package_path, exist_ok=True)
            for d in deliver_items[project][mode]:
                hal_info = hal_deliverable_info(d)
                if "架构设计" == d:
                    design_path = design_repo_path(project, mode_name)
                    if package_title == "Change":
                        change_pinout_contents[os.path.normcase(os.path.normpath(mode_package_path))] = pinout_change_content_from_latest_commits(design_path)
                    copy_directory_contents(design_path, mode_package_path)
                    if mode_name.casefold() in GPIO_MODE_NAMES:
                        for root, _, files in os.walk(mode_package_path):
                            for file in files:
                                if "引脚" in file and file.lower().endswith(".xlsx"):
                                    file_path = os.path.join(root, file)
                                    os.remove(file_path)
                                    print(f"GPIO_PIN_XLSX_REMOVED={file_path}")
                    check_middleware_filenames(mode_package_path, project, mode_name, missing_items)
                    for file_type in missing_design_file_types(mode_package_path):
                        record_missing(missing_items, project, mode_name, f"架构设计:{file_type}")
                    check_design_md5_files(mode_package_path, project, mode_name, missing_items)
                elif d.startswith("架构设计@"):
                    copy_architecture_branch(project, mode_name, d.split("@", 1)[1], mode_package_path, missing_items)
                elif d in ("main.ftl", "libopt.ftl"):
                    ftl_path = os.path.join(origin_path, project, mode_name, "ftl")
                    if not os.path.isdir(ftl_path):
                        ftl_path = next((
                            path for path in glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project), "*", "*"))
                            if os.path.isdir(path) and os.path.basename(path).casefold() == "ftl"
                        ), ftl_path)
                    for f in glob.glob(os.path.join(ftl_path, f"*{d}")):
                        shutil.copy2(f, common_path)
                    if not glob.glob(os.path.join(glob.escape(common_path), f"*{d}")):
                        record_missing(missing_items, project, mode_name, d)
                elif "PeriConfig.xml" == d:
                    peri = glob.glob(os.path.join(origin_path, project, mode_name, "peri*","PeriConfig.xml"))
                    for f in peri:
                        shutil.copy2(f, common_path)
                elif d.casefold() == PERIPHERAL_XML_FIELD:
                    copy_peripheral_xml_files(
                        project,
                        mode_name,
                        os.path.join(origin_path, f"[{project}]XML Files {datetime.now().strftime('%Y%m%d')}"),
                        missing_items,
                    )
                elif "pin" in d:
                    pin_path = os.path.join(origin_path, project, mode_name, "pin_function_list")
                    if not os.path.isdir(pin_path):
                        pin_path = next((
                            path for path in glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project), "*", "*"))
                            if os.path.isdir(path) and os.path.basename(path).casefold() == "pin_function_list"
                        ), pin_path)
                    pin_files = glob.glob(
                        os.path.join(glob.escape(pin_path), "**", f"*pin_function_list*{XLSX_SUFFIX_GLOB}"),
                        recursive=True,
                    )
                    is_gpio = mode_name.casefold() in GPIO_MODE_NAMES
                    pin_targets = [mode_package_path if is_gpio else common_path]
                    if pin_files:
                        for pin_target in pin_targets:
                            copy_files(pin_path, pin_target)
                    if not all(glob.glob(os.path.join(glob.escape(pin_target), f"*pin_function_list*{XLSX_SUFFIX_GLOB}")) for pin_target in pin_targets):
                        record_missing(missing_items, project, mode_name, "pin_function_list")
                elif series_requirement_keyword(d):
                    requirement_keyword = series_requirement_keyword(d)
                    source_files = [
                        path
                        for path in glob.glob(os.path.join(origin_path, project, mode_name, f"*{XLSX_SUFFIX_GLOB}"))
                        if series_requirement_keyword(os.path.basename(path)) == requirement_keyword
                    ]
                    if requirement_keyword == "未实现备注":
                        target_path = os.path.join(origin_path, root_path, "未实现备注")
                    else:
                        target_path = os.path.join(origin_path, f"[{project}]Common Files {datetime.now().strftime('%Y%m%d')}")
                    os.makedirs(target_path, exist_ok=True)
                    for source_file in source_files:
                        shutil.copy2(source_file, target_path)
                    if not source_files:
                        record_missing(missing_items, project, mode_name, requirement_keyword)
                elif re.fullmatch(rf"{re.escape(code_series)}_C[A-Z0-9]+", d.strip(), re.IGNORECASE):
                    project_name = case_mapping_table.get(project.lower(), "")
                    if project_name:
                        case_files = glob.glob(os.path.join(origin_path, project, mode_name, project_name, f"*{mode_name.upper()}*"))
                        for f in case_files:
                            copy_directory_contents(f, os.path.join(mode_package_path, f"{delivery_mode_name(mode_name)}_Case"))
                elif hal_info:
                    copy_hal_files(project, mode_name, d, hal_info, mode_package_path, missing_items)
                elif d.strip().upper() in LIBRARY_FOLDERS:
                    copy_hal_library(project, mode_name, d, mode_package_path, missing_items)

        os.chdir(os.path.join(origin_path, root_path))
        os.mkdir("非架构设计文件")
        os.chdir("非架构设计文件")
        for mode in sw_modes:
            mode_name = mode.rsplit("_", 1)[0]
            non_arch_mode = "Common" if mode_name == "Common" else delivery_mode_name(mode_name)
            os.mkdir(non_arch_mode)
            for d in deliver_items[project][mode]:
                target_path = os.path.join(origin_path, root_path, "非架构设计文件", non_arch_mode)
                if "反馈单" == d:
                    feedback = glob.glob(f"{origin_path}/{project}/{mode_name}/*不合格反馈单*{XLSX_SUFFIX_GLOB}")
                    for f in feedback:
                        shutil.copy2(f, target_path)
                    if not glob.glob(os.path.join(glob.escape(target_path), f"*不合格反馈单*{XLSX_SUFFIX_GLOB}")):
                        record_missing(missing_items, project, mode_name, "不合格反馈单")
                elif "变更申请" in d:
                    source_path = os.path.join(origin_path, project, mode_name)
                    feedback = glob.glob(
                        os.path.join(glob.escape(source_path), f"*变更申请单*{XLSX_SUFFIX_GLOB}")
                    )
                    feedback.extend(glob.glob(
                        os.path.join(glob.escape(source_path), "*", f"*变更申请单*{XLSX_SUFFIX_GLOB}")
                    ))
                    for f in feedback:
                        relative_dir = os.path.relpath(os.path.dirname(f), source_path)
                        copy_target = target_path if relative_dir == "." else os.path.join(target_path, relative_dir)
                        os.makedirs(copy_target, exist_ok=True)
                        shutil.copy2(f, copy_target)
                    if not feedback:
                        record_missing(missing_items, project, mode_name, "变更申请单")
                elif d == "引脚功能缺失清单":
                    unused_pin_files = glob.glob(
                        os.path.join(origin_path, project, mode_name, f"*引脚功能缺失清单*{XLSX_SUFFIX_GLOB}")
                    )
                    for f in unused_pin_files:
                        shutil.copy2(f, target_path)
                    if not glob.glob(os.path.join(glob.escape(target_path), f"*引脚功能缺失清单*{XLSX_SUFFIX_GLOB}")):
                        record_missing(missing_items, project, mode_name, "引脚功能缺失清单")
                elif "测试报告" == d:
                    continue
                elif "功能分析表" in d:
                    analysis_files = glob.glob(f"{origin_path}/{project}/{mode_name}/*外设功能分析表*{XLSX_SUFFIX_GLOB}")
                    for f in analysis_files:
                        shutil.copy2(f, target_path)
                    if not glob.glob(os.path.join(glob.escape(target_path), f"*外设功能分析表*{XLSX_SUFFIX_GLOB}")):
                        record_missing(missing_items, project, mode_name, "外设功能分析表")
                elif "tips" == d.lower():
                    tips_paths = [
                        path
                        for path in glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project), "*", "[Tt][Ii][Pp][Ss]"))
                        if os.path.isdir(path)
                    ]
                    if tips_paths:
                        copy_directory_contents(tips_paths[0], os.path.join(origin_path, root_path, "tips"))
        os.chdir(os.path.join(origin_path, root_path))
        delete_empty_folder(os.path.join(origin_path, root_path))
        # 交付清单改为读取最终打包目录，这里把实际生成的 SW 目录交给后续清单生成逻辑。
        package_paths.append(os.path.join(origin_path, root_path))
    return package_paths

# 打包交付物_AE
def packaged_deliverables_AE(deliver_items:dict, missing_items:list, package_title="Deliver List"):
    package_paths = []
    for project in deliver_items:
        code_series = re.sub(rf"^{re.escape(CONFIG['naming']['series_prefix'])}", "", project, flags=re.IGNORECASE)
        ae_modes = []
        for mode in deliver_items[project]:
            if mode.endswith("_AE"):
                ae_modes.append(mode)
        if not ae_modes:
            continue

        os.chdir(origin_path)
        root_path = f"[{project}]{package_title} To AE {datetime.now().strftime('%Y%m%d')}"
        os.mkdir(root_path)
        os.chdir(root_path)
        os.mkdir("Common")
        common_path = os.path.join(origin_path, root_path, "Common")

        for mode in ae_modes:
            mode_name = mode.rsplit("_", 1)[0]
            mode_package_path = os.path.join(origin_path, root_path, delivery_mode_name(mode_name))
            os.makedirs(mode_package_path, exist_ok=True)
            for d in deliver_items[project][mode]:
                hal_info = hal_deliverable_info(d)
                if "架构设计" == d:
                    design_path = design_repo_path(project, mode_name)
                    if package_title == "Change":
                        change_pinout_contents[os.path.normcase(os.path.normpath(mode_package_path))] = pinout_change_content_from_latest_commits(design_path)
                    copy_directory_contents(design_path, mode_package_path)
                    if mode_name.casefold() in GPIO_MODE_NAMES:
                        for root, _, files in os.walk(mode_package_path):
                            for file in files:
                                if "引脚" in file and file.lower().endswith(".xlsx"):
                                    file_path = os.path.join(root, file)
                                    os.remove(file_path)
                                    print(f"GPIO_PIN_XLSX_REMOVED={file_path}")
                    check_middleware_filenames(mode_package_path, project, mode_name, missing_items)
                    for file_type in missing_design_file_types(mode_package_path):
                        record_missing(missing_items, project, mode_name, f"架构设计:{file_type}")
                    check_design_md5_files(mode_package_path, project, mode_name, missing_items)
                    reference_path = design_reference_paths.get(
                        (project.casefold(), mode_name.casefold()), ""
                    ) if package_title == "Deliver List" else ""
                    if reference_path:
                        copy_directory_contents(reference_path, os.path.join(mode_package_path, "reference_version"))
                elif d.startswith("架构设计@"):
                    copy_architecture_branch(project, mode_name, d.split("@", 1)[1], mode_package_path, missing_items)
                elif d in ("main.ftl", "libopt.ftl"):
                    ftl_path = os.path.join(origin_path, project, mode_name, "ftl")
                    if not os.path.isdir(ftl_path):
                        ftl_path = next((
                            path for path in glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project), "*", "*"))
                            if os.path.isdir(path) and os.path.basename(path).casefold() == "ftl"
                        ), ftl_path)
                    for f in glob.glob(os.path.join(ftl_path, f"*{d}")):
                        shutil.copy2(f, common_path)
                    if not glob.glob(os.path.join(glob.escape(common_path), f"*{d}")):
                        record_missing(missing_items, project, mode_name, d)
                elif "PeriConfig.xml" == d:
                    peri = glob.glob(os.path.join(origin_path, project, mode_name, "peri*","PeriConfig.xml"))
                    for f in peri:
                        shutil.copy2(f, os.path.join(origin_path, root_path, "Common"))
                        shutil.copy2(f, common_path)
                elif d.casefold() == PERIPHERAL_XML_FIELD:
                    copy_peripheral_xml_files(
                        project,
                        mode_name,
                        os.path.join(origin_path, f"[{project}]XML Files {datetime.now().strftime('%Y%m%d')}"),
                        missing_items,
                    )
                elif "pin" in d:
                    pin_path = os.path.join(origin_path, project, mode_name, "pin_function_list")
                    if not os.path.isdir(pin_path):
                        pin_path = next((
                            path for path in glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project), "*", "*"))
                            if os.path.isdir(path) and os.path.basename(path).casefold() == "pin_function_list"
                        ), pin_path)
                    pin_files = glob.glob(
                        os.path.join(glob.escape(pin_path), "**", f"*pin_function_list*{XLSX_SUFFIX_GLOB}"),
                        recursive=True,
                    )
                    is_gpio = mode_name.casefold() in GPIO_MODE_NAMES
                    pin_targets = [mode_package_path if is_gpio else common_path]
                    if pin_files:
                        for pin_target in pin_targets:
                            copy_files(pin_path, pin_target)
                    if not all(glob.glob(os.path.join(glob.escape(pin_target), f"*pin_function_list*{XLSX_SUFFIX_GLOB}")) for pin_target in pin_targets):
                        record_missing(missing_items, project, mode_name, "pin_function_list")
                elif series_requirement_keyword(d):
                    requirement_keyword = series_requirement_keyword(d)
                    source_files = [
                        path
                        for path in glob.glob(os.path.join(origin_path, project, mode_name, f"*{XLSX_SUFFIX_GLOB}"))
                        if series_requirement_keyword(os.path.basename(path)) == requirement_keyword
                    ]
                    target_path = os.path.join(origin_path, f"[{project}]Common Files {datetime.now().strftime('%Y%m%d')}")
                    os.makedirs(target_path, exist_ok=True)
                    for source_file in source_files:
                        shutil.copy2(source_file, target_path)
                    if not source_files:
                        record_missing(missing_items, project, mode_name, requirement_keyword)
                elif "评审表" == d:
                    feedback = [
                        f
                        for f in glob.glob(os.path.join(origin_path, project, mode_name, f"*{XLSX_SUFFIX_GLOB}"))
                        if report_type(os.path.basename(f)) == "review"
                    ]
                    report_path = mode_package_path
                    if mode_name == "Common":
                        report_path = os.path.join(origin_path, root_path, "非架构设计文件", "Common")
                        os.makedirs(report_path, exist_ok=True)
                    for f in feedback:
                        shutil.copy2(f, report_path)
                    if not any(
                        report_type(os.path.basename(path)) == "review"
                        for path in glob.glob(os.path.join(glob.escape(report_path), f"*{XLSX_SUFFIX_GLOB}"))
                    ):
                        record_missing(missing_items, project, mode_name, "评审表")
                elif "测试报告" == d:
                    reports = [
                        f
                        for f in glob.glob(os.path.join(origin_path, project, mode_name, f"*{XLS_SUFFIX_GLOB}"))
                        if report_type(os.path.basename(f)) == "test_report"
                    ]
                    report_path = mode_package_path
                    if mode_name == "Common":
                        report_path = os.path.join(origin_path, root_path, "非架构设计文件", "Common")
                        os.makedirs(report_path, exist_ok=True)
                    for f in reports:
                        shutil.copy2(f, report_path)
                    if not any(
                        report_type(os.path.basename(path)) == "test_report"
                        for path in glob.glob(os.path.join(glob.escape(report_path), f"*{XLS_SUFFIX_GLOB}"))
                    ):
                        record_missing(missing_items, project, mode_name, "测试报告")
                elif "功能分析表" in d:
                    analysis_files = glob.glob(os.path.join(origin_path, project, mode_name, f"*外设功能分析表*{XLSX_SUFFIX_GLOB}"))
                    analysis_path = os.path.join(origin_path, root_path, "外设功能分析表", delivery_mode_name(mode_name))
                    os.makedirs(analysis_path, exist_ok=True)
                    for f in analysis_files:
                        shutil.copy2(f, analysis_path)
                    if not glob.glob(os.path.join(glob.escape(analysis_path), f"*外设功能分析表*{XLSX_SUFFIX_GLOB}")):
                        record_missing(missing_items, project, mode_name, "外设功能分析表")
                elif "tips" == d.lower():
                    tips_paths = [
                        path
                        for path in glob.glob(os.path.join(glob.escape(origin_path), glob.escape(project), "*", "[Tt][Ii][Pp][Ss]"))
                        if os.path.isdir(path)
                    ]
                    if tips_paths:
                        copy_directory_contents(tips_paths[0], os.path.join(origin_path, root_path, "tips"))
                elif re.fullmatch(rf"{re.escape(code_series)}_C[A-Z0-9]+", d.strip(), re.IGNORECASE):
                    project_name = case_mapping_table.get(project.lower(), "")
                    if project_name:
                        case_files = glob.glob(os.path.join(origin_path, project, mode_name, project_name, f"*{mode_name.upper()}*"))
                        for f in case_files:
                            copy_directory_contents(f, os.path.join(mode_package_path, f"{delivery_mode_name(mode_name)}_Case"))
                elif hal_info:
                    copy_hal_files(project, mode_name, d, hal_info, mode_package_path, missing_items)
                elif d.strip().upper() in LIBRARY_FOLDERS:
                    copy_hal_library(project, mode_name, d, mode_package_path, missing_items)
        delete_empty_folder(os.path.join(origin_path, root_path))
        # 交付清单改为读取最终打包目录，这里把实际生成的 AE 目录交给后续清单生成逻辑。
        package_paths.append(os.path.join(origin_path, root_path))
    return package_paths

def next_version_path(path):
    if not os.path.exists(path):
        return path

    root, ext = os.path.splitext(path)
    index = 1
    while os.path.exists(f"{root}-{index}{ext}"):
        index += 1
    return f"{root}-{index}{ext}"


def append_cell_value(row, column, value):
    if row[column]:
        row[column] += "\n" + value
    else:
        row[column] = value


def list_deliver_files(path):
    # 交付物清单必须来自实际拉下来的文件
    if not os.path.exists(path):
        return []

    files = []
    ignored_dirs = {".git"} | module_dir_blacklist()
    for root, dirs, names in os.walk(path):
        root_dirs = {part.casefold() for part in Path(root).parts}
        if root_dirs & ignored_dirs:
            dirs[:] = []
            continue
        dirs[:] = [d for d in dirs if d.casefold() not in ignored_dirs]
        for name in names:
            if name.startswith(".") or name.startswith("~$"):
                continue
            ext = os.path.splitext(name)[1].casefold()
            if is_excel_filename(name) or ext in [".xml", ".ftl", ".txt"]:
                # if ext == ".xml" and root_dirs & ignored_dirs:
                #     files.append(f"其他:{name}")
                # else:
                #     files.append(name)
                files.append(name)
    return sorted(set(files))


def standard_driver_versions(path):
    # 标准驱动列取 .c/.h 文件开头 version 后的日期；每个文件只读前 500 字符。
    if not os.path.exists(path):
        return []

    versions = []
    for root, dirs, names in os.walk(path):
        dirs[:] = [d for d in dirs if d != ".git"]
        for name in names:
            if os.path.splitext(name)[1].lower() not in [".c", ".h"]:
                continue
            path = os.path.join(root, name)
            with open(path, "r", encoding="utf-8", errors="ignore") as file:
                content = file.read(500)
            match = re.search(r"version\s+(\d{4}[-/]\d{1,2}[-/]\d{1,2})", content, re.IGNORECASE)
            if match:
                versions.append(match.group(1).replace("/", "-"))
    return sorted(set(versions))


def package_driver_versions(package_path, module_path):
    current_path = module_path
    while os.path.normcase(current_path) != os.path.normcase(package_path):
        versions = standard_driver_versions(current_path)
        if versions:
            return versions
        current_path = os.path.dirname(current_path)
    return []

def module_list_config():
    global module_list_config_cache
    if module_list_config_cache is not None:
        return module_list_config_cache

    config = {
        "blacklist": {item.casefold() for item in CONFIG["module_list"]["blacklist"]},
        "whitelist": {item.casefold() for item in CONFIG["module_list"]["whitelist"]},
    }
    issue_id = CONFIG["module_list"].get("issue_id")
    if not issue_id:
        module_list_config_cache = config
        return config
    config_issue = redmine.issue.get(issue_id)

    current_section = ""
    for raw_line in str(getattr(config_issue, "description", "") or "").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        section = line.strip("[]").strip().casefold() if line.startswith("[") and line.endswith("]") else ""
        if section in ["blacklist", "black", "黑名单"]:
            current_section = "blacklist"
            continue
        if section in ["whitelist", "white", "白名单"]:
            current_section = "whitelist"
            continue
        if current_section:
            config[current_section].add(line.casefold())

    print(f"MODULE_LIST_CONFIG={config_issue.id}\t{config_issue.subject}")
    module_list_config_cache = config
    return module_list_config_cache

# 黑名单目录只用于交付清单模块识别，不影响实际打包复制。
def module_dir_blacklist():
    return module_list_config()["blacklist"]

# 白名单目录仅影响交付清单中模块名称的显示，不影响实际文件和目录的识别与复制。
def module_prefix_whitelist():
    return module_list_config()["whitelist"]


def deliver_list_modules(project, mode):
    design_path = design_repo_path(project, mode)
    if not os.path.isdir(design_path):
        # 没有架构设计实例目录时不写交付清单行。
        return []

    module_dirs = sorted(
        (
            os.path.join(design_path, name)
            for name in module_dir_names(design_path)
        ),
        key=lambda path: os.path.basename(path).casefold(),
    )
    if not module_dirs:
        return [(mode, design_path, design_path)]

    return [(os.path.basename(path), path, design_path) for path in module_dirs]

def change_application_files(project, mode):
    # 变更申请单存放在 mode 外层目录，需要单独补进交付清单“其他”列。
    return sorted(
        os.path.basename(path)
        for path in glob.glob(os.path.join(origin_path, project, mode, f"*变更申请单*{XLSX_SUFFIX_GLOB}"))
    )


#生成common交付物列表，供交付清单和打包使用；common 目录下的文件不要求必须存在，缺失时记录但不阻断流程。
def common_deliver_files(project, modes):
    files = set()
    for mode in modes:
        common_patterns = [
            os.path.join(origin_path, project, mode, "ftl", "*main.ftl"),
            os.path.join(origin_path, project, mode, "ftl", "*libopt.ftl"),
            os.path.join(origin_path, project, mode, "pin_function_list", f"*pin_function_list*{XLSX_SUFFIX_GLOB}"),
            os.path.join(origin_path, project, mode, "peri*", "PeriConfig.xml"),
            os.path.join(origin_path, project, mode, "peripheral_xml", "*.xml"),
            os.path.join(origin_path, project, mode, "perihearl_xml", "*.xml"),
            os.path.join(origin_path, project, mode, "xml", "*.xml"),
            os.path.join(origin_path, project, mode, "peripheral_config_information", "*.xml"),
        ]
        for pattern in common_patterns:
            for path in glob.glob(pattern):
                files.add(os.path.basename(path))
    return sorted(files)


def package_module_dirs(package_path):
    ignored_dirs = {".git", "branches", "common", "tips", "reference_version", "peripheral_config_information", "非架构设计文件".casefold(), "备注"} | module_dir_blacklist()
    whitelist = module_prefix_whitelist() - ignored_dirs
    module_file_exts = {".xml", ".ftl", ".txt", ".c", ".h"}
    module_dirs = []
    branch_parent_dirs = []
    for root, dirs, files in os.walk(package_path):
        rel_parts = Path(os.path.relpath(root, package_path)).parts
        rel_keys = [part.casefold() for part in rel_parts]
        if any(part in ignored_dirs for part in rel_keys):
            dirs[:] = []
            continue
        if any(name.casefold() == "branches" for name in dirs):
            branch_parent_dirs.append(root)
        example_dir = f"{rel_parts[0]}_example".casefold() if len(rel_parts) == 1 else ""
        dirs[:] = [
            name
            for name in dirs
            if name.casefold() not in ignored_dirs and name != ".git" and name.casefold() != example_dir
        ]
        if root == package_path:
            continue
        if os.path.basename(root).casefold() in whitelist:
            # 白名单目录本身就是交付清单模块，命中后不再把 Include/Source 等子目录当模块。
            module_dirs.append(root)
            dirs[:] = []
            continue
        if any(
            not name.startswith(".")
            and not name.startswith("~$")
            and (
                is_excel_filename(name)
                or os.path.splitext(name)[1].casefold() in module_file_exts
            )
            for name in files
        ):
            module_dirs.append(root)

    # 只有附加分支内容时，默认分支没有文件可供普通规则识别，使用分支父目录作为模块行。
    for parent in branch_parent_dirs:
        if not any(
            path == parent or os.path.commonpath([parent, path]) == parent
            for path in module_dirs
        ):
            module_dirs.append(parent)

    # 有些仓库会多套一层目录，清单只取最具体的交付目录名，避免把外层目录误当模块。
    return sorted(
        (
            path
            for path in module_dirs
            if not any(
                other != path and os.path.commonpath([path, other]) == path
                for other in module_dirs
            )
        ),
        key=lambda path: os.path.relpath(path, package_path).casefold(),
    )


def package_extra_files(package_path, module_path):
    # 非架构设计文件已经在打包阶段按模块归档；生成表格时统一并入对应模块的“其他”列。
    rel_parts = Path(os.path.relpath(module_path, package_path)).parts
    extra_modules = {rel_parts[0].upper(), rel_parts[-1].upper()}
    files = []
    for module in extra_modules:
        extra_path = os.path.join(package_path, "非架构设计文件", module)
        files.extend(f"其他:{name}" for name in list_deliver_files(extra_path))
    return sorted(set(files))


def package_module_name(package_path, module_path):
    rel_parts = Path(os.path.relpath(module_path, package_path)).parts
    blacklist = module_dir_blacklist()
    whitelist = module_prefix_whitelist() - blacklist
    rel_keys = [part.casefold() for part in rel_parts]
    if len(rel_parts) > 1 and not any(part in blacklist for part in rel_keys) and rel_keys[-1] in whitelist:
        # 白名单外层目录只作为模块前缀显示，实际文件和目录不改名。
        return f"{rel_parts[0]}_{rel_parts[-1]}".upper()
    return rel_parts[-1].upper()


def is_interface_table(filename):
    if filename.startswith("~$"):
        return False
    lower_name = filename.lower()
    if not is_excel_filename(filename):
        return False
    # 客户反馈类报告也是 Excel 文件，但不属于 IDE 界面表模板列。
    if any(key in filename for key in ["开发环节评审表", "不合格反馈单"]):
        return False
    return "界面表" in filename or "embedded builder" in lower_name


def cell_text(value):
    # 统一把 Excel 单元格值转成字符串，空值按空串处理，便于后面对比引脚功能。
    if value in [None, ""]:
        return ""
    if isinstance(value, float) and value.is_integer():
        # openpyxl 可能把整数单元格读成 1.0，这里避免差异内容里多出无意义的小数。
        return str(int(value))
    return str(value).strip()


def interface_table_history_key(path):
    # 界面表文件名会被脚本按 content history 追加 _v版本号。
    # 用 git 历史比较时先去掉版本后缀，避免同一张表因为版本号变化被误判为新增/删除。
    path_parts = path.replace("\\", "/").split("/")
    stem = re.sub(r"_v\d+(?:\.\d+)*$", "", Path(path_parts[-1]).stem, flags=re.IGNORECASE)
    path_parts[-1] = f"{stem}.xlsx"
    return "/".join(part.casefold() for part in path_parts)


def summarize_pin_prefix_change(old_pin, new_pin):
    old_parts = [part.strip() for part in re.split(r"[,，]", old_pin) if part.strip()]
    new_parts = [part.strip() for part in re.split(r"[,，]", new_pin) if part.strip()]
    if len(old_parts) != len(new_parts) or not old_parts:
        return old_pin, new_pin

    old_prefixes = set()
    new_prefixes = set()
    old_suffixes = []
    new_suffixes = []
    for old_part, new_part in zip(old_parts, new_parts):
        if "_" not in old_part or "_" not in new_part:
            return old_pin, new_pin
        old_prefix, old_suffix = old_part.split("_", 1)
        new_prefix, new_suffix = new_part.split("_", 1)
        old_prefixes.add(old_prefix)
        new_prefixes.add(new_prefix)
        old_suffixes.append(old_suffix)
        new_suffixes.append(new_suffix)

    if old_suffixes == new_suffixes and len(old_prefixes) == 1 and len(new_prefixes) == 1 and old_prefixes != new_prefixes:
        return f"{old_prefixes.pop()}_", f"{new_prefixes.pop()}_"
    return old_pin, new_pin


def pinout_rows_from_workbook_blob(blob):
    workbook = load_workbook(BytesIO(blob), read_only=True, data_only=True)
    try:
        sheet = next((current_sheet for current_sheet in workbook.worksheets if current_sheet.title.strip().casefold() == "pinout"), None)
        if sheet is None:
            return {}

        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header_indexes = {
            cell_text(value): index
            for index, value in enumerate(header_values)
            if cell_text(value)
        }
        pin_function_index = header_indexes.get("引脚功能")
        if pin_function_index is None:
            return {}

        filled_key_indexes = [
            header_indexes[name]
            for name in ["外设", "组名称", "功能名称"]
            if name in header_indexes
        ]
        sub_item_index = header_indexes.get("子项名称")
        filled_values = {index: "" for index in filled_key_indexes}
        pinout_rows = {}

        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if values and cell_text(values[0]).casefold() == "end":
                break

            for index in filled_key_indexes:
                value = cell_text(values[index]) if index < len(values) else ""
                if value:
                    filled_values[index] = value

            pin_function = cell_text(values[pin_function_index]) if pin_function_index < len(values) else ""
            if pin_function:
                # 同一个“引脚功能”单元格里可能重复列同一个引脚名，这里只去掉单元格内重复项。
                # 例如 ETH_MDIO, ETH_MDC, ETH_MDIO 会写成 ETH_MDIO, ETH_MDC。
                pin_parts = []
                for part in re.split(r"[,，]", pin_function):
                    part = part.strip()
                    if part and part not in pin_parts:
                        pin_parts.append(part)
                if pin_parts:
                    pin_function = ", ".join(pin_parts)
            sub_item = cell_text(values[sub_item_index]) if sub_item_index is not None and sub_item_index < len(values) else ""
            key_parts = [filled_values[index] for index in filled_key_indexes] + [sub_item]
            display_parts = [part for part in key_parts if part and part != "/"]
            if not display_parts and not pin_function:
                continue

            key = tuple(key_parts) if any(key_parts) else (f"row:{row_number}",)
            pinout_rows[key] = (" / ".join(display_parts) or f"第{row_number}行", pin_function)

        return pinout_rows
    finally:
        workbook.close()


def pinout_change_content_from_latest_commits(repo_path):
    # Change 包的“修改内容”不读取 Redmine 的 commit id，避免字段填错影响结果。
    # 这里直接比较架构设计仓库当前最新提交 HEAD 和上一次提交 HEAD^ 中的界面表。
    if not os.path.isdir(os.path.join(repo_path, ".git")):
        return ""

    try:
        head_commit = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=repo_path,
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        ).stdout.strip()
        previous_commit = subprocess.run(
            ["git", "rev-parse", "HEAD^"],
            cwd=repo_path,
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        ).stdout.strip()
        print(f"PINOUT_CHANGE_COMPARE={repo_path}\told={previous_commit}\tnew={head_commit}")
        table_paths_by_commit = {}
        for commit in [previous_commit, head_commit]:
            result = subprocess.run(
                ["git", "-c", "core.quotePath=false", "ls-tree", "-r", "--name-only", commit],
                cwd=repo_path,
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
            table_paths_by_commit[commit] = {
                interface_table_history_key(path): path
                for path in result.stdout.splitlines()
                if is_interface_table(os.path.basename(path))
            }
    except subprocess.CalledProcessError as error:
        print(f"PINOUT_CHANGE_DIFF_FAILED={repo_path}\t{error}")
        return ""

    old_tables = table_paths_by_commit[previous_commit]
    new_tables = table_paths_by_commit[head_commit]
    change_lines = []

    for table_key in sorted(set(old_tables) | set(new_tables)):
        old_path = old_tables.get(table_key)
        new_path = new_tables.get(table_key)
        table_name = os.path.basename(new_path or old_path)
        print(f"PINOUT_CHANGE_TABLE={table_name}\told={old_path or '无'}\tnew={new_path or '无'}")
        try:
            old_rows = {}
            new_rows = {}
            if old_path:
                old_blob = subprocess.run(
                    ["git", "show", f"{previous_commit}:{old_path}"],
                    cwd=repo_path,
                    check=True,
                    capture_output=True,
                ).stdout
                old_rows = pinout_rows_from_workbook_blob(old_blob)
            if new_path:
                new_blob = subprocess.run(
                    ["git", "show", f"{head_commit}:{new_path}"],
                    cwd=repo_path,
                    check=True,
                    capture_output=True,
                ).stdout
                new_rows = pinout_rows_from_workbook_blob(new_blob)
        except Exception as error:
            print(f"PINOUT_CHANGE_READ_FAILED={repo_path}\t{table_name}\t{error}")
            continue

        table_change_count = 0
        for row_key in sorted(set(old_rows) | set(new_rows), key=lambda item: " / ".join(item)):
            old_display, old_pin = old_rows.get(row_key, (new_rows.get(row_key, ("", ""))[0], ""))
            new_display, new_pin = new_rows.get(row_key, (old_display, ""))
            display = new_display or old_display
            if old_pin == new_pin:
                continue
            table_change_count += 1
            if old_pin and new_pin:
                old_pin, new_pin = summarize_pin_prefix_change(old_pin, new_pin)
                change_lines.append(f"{table_name}：{display} 引脚功能由 {old_pin} 改为 {new_pin}")
            elif new_pin:
                change_lines.append(f"{table_name}：{display} 新增引脚功能 {new_pin}")
            else:
                change_lines.append(f"{table_name}：{display} 删除引脚功能 {old_pin}")
        if old_rows or new_rows:
            print(f"PINOUT_CHANGE_RESULT={table_name}\tchanges={table_change_count}")
        else:
            print(f"PINOUT_CHANGE_RESULT={table_name}\t无可比对pinout内容")

    return "\n".join(change_lines)


def content_history_version(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        # 新模板优先使用 content history；旧模板没有该 sheet 或没有版本记录时，兼容读取 history。
        sheet_by_name = {
            current_sheet.title.strip().casefold(): current_sheet
            for current_sheet in workbook.worksheets
        }
        latest_version_cell = None
        for sheet_name in ("content history", "history"):
            sheet = sheet_by_name.get(sheet_name)
            if sheet is None:
                continue
            # 部分 WPS 文件的 worksheet dimension 错写为 A1:A1，只读模式下需重置后才能读到后续版本行。
            sheet.reset_dimensions()
            # 版本表 A 列最后一个非空单元格是最新版本号，首行“文件版本”等表头不参与判断。
            for (cell,) in sheet.iter_rows(min_row=2, min_col=1, max_col=1):
                if cell.value not in [None, ""]:
                    latest_version_cell = cell
            if latest_version_cell is not None:
                break
        if latest_version_cell is None:
            return ""

        value = latest_version_cell.value
        if isinstance(value, (int, float)):
            decimal_format = re.search(r"0\.(0+)", latest_version_cell.number_format)
            version = f"{value:.{len(decimal_format.group(1))}f}" if decimal_format else str(value)
        else:
            version = str(value).strip()
        return version[1:] if version[:1].casefold() == "v" else version
    finally:
        workbook.close()


def rename_interface_tables(repository_path):
    # 界面表和中间件工程配置 Excel 按 content history/history 的最新版本重命名，已有版本后缀会被替换。
    for root, dirs, files in os.walk(repository_path):
        dirs[:] = [directory for directory in dirs if directory != ".git"]
        for filename in files:
            lower_name = filename.lower()
            if filename.startswith("~$") or not lower_name.endswith(".xlsx"):
                continue
            is_versioned_excel = is_interface_table(filename) or "工程配置" in filename
            if not is_versioned_excel:
                continue

            source_path = Path(root) / filename
            version = content_history_version(source_path)
            if not version:
                # 这些 Excel 是对外交付文件，没有版本号时直接失败，避免交付出未标版本的文件。
                raise RuntimeError(f"xlsx缺少content history/history版本号: {source_path}")

            stem = re.sub(r"_v\d+(?:\.\d+)*$", "", source_path.stem, flags=re.IGNORECASE)
            target_path = source_path.with_name(f"{stem}_v{version}{source_path.suffix}")
            if target_path == source_path:
                continue
            if target_path.exists():
                raise FileExistsError(f"xlsx版本文件已存在: {target_path}")

            source_path.rename(target_path)
            log_key = "INTERFACE_TABLE_RENAMED" if is_interface_table(filename) else "XLSX_RENAMED"
            print(f"{log_key}={source_path.name}\t{target_path.name}")


def rename_config_xml_files(repository_path):
    # 配置 xml 的版本号写在根节点 FileVersion 中，实际文件名也要补上 _v版本号。
    for root, dirs, files in os.walk(repository_path):
        dirs[:] = [directory for directory in dirs if directory != ".git"]
        for filename in files:
            if not is_config_xml(filename):
                continue

            source_path = Path(root) / filename
            # 示例：<Root FileVersion="1.0" Version="V2.0">，交付命名只取 FileVersion。
            version = str(ET.parse(source_path).getroot().get("FileVersion", "")).strip()
            if not version:
                continue

            stem = re.sub(r"_v\d+(?:\.\d+)*$", "", source_path.stem, flags=re.IGNORECASE)
            target_path = source_path.with_name(f"{stem}_v{version}{source_path.suffix}")
            if target_path == source_path:
                continue
            if target_path.exists():
                raise FileExistsError(f"配置xml版本文件已存在: {target_path}")

            source_path.rename(target_path)
            print(f"CONFIG_XML_RENAMED={source_path.name}\t{target_path.name}")


def is_config_xml(filename):
    lower_name = filename.lower()
    return lower_name.endswith(".xml") and ("config" in lower_name or "配置" in filename)


def build_deliver_list_row(mode, deliverables, files, driver_versions):
    # 按文档中的交付物清单列归类；无法明确归类的 Excel/xml/ftl 放入“其他”列。
    row = {
        "Peripheral": mode.upper(),
        "IDE界面表模板": "",
        "标准驱动": "",
        "xx_config.xml": "",
        "MD5": "-",
        "其他": "",
    }

    for version in driver_versions:
        append_cell_value(row, "标准驱动", version)

    for item in deliverables:
        if "md5" in item.lower():
            row["MD5"] = "YES"

    for filename in files:
        if filename.startswith("其他:"):
            append_cell_value(row, "其他", filename.removeprefix("其他:"))
            continue
        lower_name = filename.lower()
        if is_interface_table(filename):
            append_cell_value(row, "IDE界面表模板", filename)
        elif is_config_xml(filename):
            append_cell_value(row, "xx_config.xml", filename)
        elif "md5" in lower_name:
            row["MD5"] = "YES"
        elif is_excel_filename(filename) or lower_name.endswith((".xml", ".ftl")):
            append_cell_value(row, "其他", filename)

    return row

# 生成交付清单 xlsx 文件
def write_deliver_list_workbook(rows, list_path, include_change_content=False):
    headers = ["Peripheral", "IDE界面表模板", "标准驱动", "xx_config.xml", "MD5", "其他"]
    if include_change_content:
        headers.append("修改内容")
    header_fill = PatternFill("solid", fgColor="B7F15A")
    thin_side = Side(style="thin", color="000000")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Deliver List"
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if cell.row == 1:
                cell.fill = header_fill
                cell.font = Font(bold=True)

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 10
    sheet.column_dimensions["F"].width = 48
    if include_change_content:
        sheet.column_dimensions["G"].width = 72

    workbook.save(list_path)


def generate_deliver_list_from_packages(package_paths):
    for package_path in package_paths:
        if not os.path.isdir(package_path):
            continue

        project_match = re.match(r"\[([^\]]+)\](?:Deliver List|Change) To ", os.path.basename(package_path))
        project = project_match.group(1) if project_match else os.path.basename(package_path)
        is_change_package = re.match(r"\[[^\]]+\]Change To ", os.path.basename(package_path)) is not None
        rows = []

        # 表格从最终打包目录反扫，保证清单内容和客户实际收到的目录一致。
        for module_path in package_module_dirs(package_path):
            module = package_module_name(package_path, module_path)
            files = list_deliver_files(module_path)
            files.extend(package_extra_files(package_path, module_path))
            row = build_deliver_list_row(module, [], files, package_driver_versions(package_path, module_path))
            if is_change_package:
                module_norm = os.path.normcase(os.path.normpath(module_path))
                row["修改内容"] = ""
                for change_root, change_content in change_pinout_contents.items():
                    try:
                        if os.path.commonpath([module_norm, change_root]) == change_root:
                            row["修改内容"] = change_content
                            break
                    except ValueError:
                        pass
            rows.append(row)
            print(f"DELIVER_ITEM={project}\t{module}\t{','.join(files)}")

        for example_path in glob.glob(os.path.join(glob.escape(package_path), "*", "*_Example")):
            module = os.path.basename(os.path.dirname(example_path)).upper()
            row = next((item for item in rows if item["Peripheral"].casefold() == module.casefold()), None)
            if row:
                append_cell_value(row, "其他", "Example")
            else:
                row = build_deliver_list_row(module, [], [], [])
                row["其他"] = "Example"
                if is_change_package:
                    row["修改内容"] = ""
                rows.append(row)
            print(f"DELIVER_ITEM={project}\t{module}\tExample")

        common_files = sorted(set(
            list_deliver_files(os.path.join(package_path, "Common"))
            + list_deliver_files(os.path.join(package_path, "peripheral_config_information"))
        ))
        if common_files:
            row = build_deliver_list_row("Common", [], common_files, [])
            if is_change_package:
                row["修改内容"] = ""
            rows.append(row)
            print(f"DELIVER_ITEM={project}\tCOMMON\t{','.join(common_files)}")

        list_path = next_version_path(f"{package_path}.xlsx")
        write_deliver_list_workbook(rows, list_path, is_change_package)
        print(f"DELIVER_LIST={list_path}")


# 生成交付物清单 xlsx，文件用于归档，print 用于 Jenkins 日志收集。
def generate_deliver_list(deliver_items:dict):
    headers = ["Peripheral", "IDE界面表模板", "标准驱动", "xx_config.xml", "MD5", "其他"]
    header_fill = PatternFill("solid", fgColor="B7F15A")
    thin_side = Side(style="thin", color="000000")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for project in sorted(deliver_items):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Deliver List"
        sheet.append(headers)

        for mode in sorted(deliver_items[project]):
            for module, module_path, driver_path in deliver_list_modules(project, mode):
                files = list_deliver_files(module_path)
                # 钉盘额外文件不在架构设计模块目录里，不能靠 list_deliver_files 扫出来。
                if any("变更申请" in item for item in deliver_items[project][mode]):
                    files.extend(change_application_files(project, mode))
                if any("功能分析表" in item for item in deliver_items[project][mode]):
                    files.extend(
                        os.path.basename(path)
                        for path in glob.glob(os.path.join(origin_path, project, mode, f"*外设功能分析表*{XLSX_SUFFIX_GLOB}"))
                    )
                driver_versions = standard_driver_versions(driver_path)
                row = build_deliver_list_row(module, sorted(deliver_items[project][mode]), files, driver_versions)
                sheet.append([row[header] for header in headers])
                print(f"DELIVER_ITEM={project}\t{module}\t{','.join(files)}")

        common_files = common_deliver_files(project, deliver_items[project])
        if common_files:
            row = build_deliver_list_row("Common", [], common_files, [])
            sheet.append([row[header] for header in headers])
            print(f"DELIVER_ITEM={project}\tCOMMON\t{','.join(common_files)}")
        #格式居中等....
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = Font(bold=True)

        sheet.column_dimensions["A"].width = 16
        sheet.column_dimensions["B"].width = 34
        sheet.column_dimensions["C"].width = 18
        sheet.column_dimensions["D"].width = 28
        sheet.column_dimensions["E"].width = 10
        sheet.column_dimensions["F"].width = 48

        project_deliverables = [d for mode in deliver_items[project] for d in deliver_items[project][mode]]
        targets = ["SW"] if any("变更申请" in d for d in project_deliverables) else ["AE"]
        for target in targets:
            list_name = f"[{project}]Deliver List To {target} {datetime.now().strftime('%Y%m%d')}.xlsx"
            list_path = next_version_path(os.path.join(origin_path, list_name))
            workbook.save(list_path)
            print(f"DELIVER_LIST={list_path}")

# TODO: 同一大系列下的小系列需要做识别与区分
# TODO: 从钉盘拉取文件
def deliver():
    # 提前创建交付根目录，避免后续 os.chdir(origin_path) 时因目录不存在失败。
    os.makedirs(origin_path, exist_ok=True)
    architecture_branch_paths.clear()
    design_reference_paths.clear()

    deliver_items = {}
    change_deliver_items = {}
    missing_items = []
    projects = get_git_url_for_repo()
    issues = redmine.issue.filter(project_id=CONFIG["redmine"]["project_id"], status_id="*")
    deliver_lists = redmine.issue.filter(project_id=CONFIG["redmine"]["project_id"], subject=CONFIG["redmine"]["delivery_subject"])
    change_milestones = redmine.issue.filter(project_id=CONFIG["redmine"]["project_id"], subject=CONFIG["redmine"]["change_subject"])
    change_issues = issues

    need_deliver = {}
    need_project_examples = set()
    need_project_designs = set()
    need_project_firmwares = set()
    # 解析交付清单，获取需要交付的外设
    for deliver in deliver_lists:
        lines = [line.strip() for line in deliver.description.splitlines() if line.strip()]
        time = datetime.strptime(lines[0], "%Y-%m-%d").date()
        device = lines[1]
        current_time = date.today()

        if time == current_time:
            if TARGET_SERIES_SET and deliver.project.name not in TARGET_SERIES_SET:
                continue
            need_deliver[deliver.project.name] = []
            for item in re.split(r"[,，]", device):
                item = item.strip()
                if not item:
                    continue
                if item.upper() == "EXAMPLE":
                    need_project_examples.add(deliver.project.name)
                elif item.upper() == "FIRMWARE":
                    need_project_firmwares.add(deliver.project.name)
                elif item == "架构设计":
                    need_project_designs.add(deliver.project.name)
                else:
                    need_deliver[deliver.project.name].append(item)

    change_parent_ids = {str(milestone.id) for milestone in change_milestones}
    regular_issue_ids = {str(issue.id) for issue in issues}
    all_issues = []
    seen_issue_ids = set()
    for issue in list(issues) + list(change_issues):
        issue_id = str(getattr(issue, "id", ""))
        if issue_id in seen_issue_ids:
            continue
        seen_issue_ids.add(issue_id)
        all_issues.append(issue)

    design_reference_commits = {}
    full_design_projects = {project.casefold() for project in need_project_designs}
    # 参考版本独立查询全部状态的问题单，不设总数上限，由客户端自动分页。
    reference_issues = redmine.issue.filter(
        project_id=CONFIG["redmine"]["project_id"], status_id='*'
    ) if full_design_projects else ()
    for issue in reference_issues:
        project_name = getattr(getattr(issue, "project", None), "name", "")
        mode = getattr(getattr(issue, "category", None), "name", "")
        if not project_name or project_name.casefold() not in full_design_projects or not mode:
            continue
        for field in issue.custom_fields:
            if field["name"] != CONFIG["redmine"]["custom_fields"]["reference_commit_id"]:
                continue
            reference_commit = str(field["value"] or "").strip()
            if reference_commit:
                design_reference_commits.setdefault(project_name.casefold(), {}).setdefault(
                    mode.casefold(), reference_commit
                )
            break

    for issue in all_issues:
        try:
            issue_id = str(getattr(issue, "id", ""))
            project_name = issue.project.name
            if TARGET_SERIES_SET and project_name not in TARGET_SERIES_SET:
                continue
            code_series = re.sub(rf"^{re.escape(CONFIG['naming']['series_prefix'])}", "", project_name, flags=re.IGNORECASE)
            dingtalk_series = CONFIG["dingtalk"]["series_aliases"].get(project_name.casefold(), project_name)
            status = issue.status.name
            # 普通交付和 Change 交付都只处理 Verified 状态的 Feature。
            parent_id = str(getattr(getattr(issue, "parent", None), "id", ""))
            category_name = getattr(getattr(issue, "category", None), "name", "")
            is_regular_deliver = issue_id in regular_issue_ids and status == CONFIG["redmine"]["status"] and project_name in need_deliver and category_name in need_deliver[project_name]
            is_change_deliver = status == CONFIG["redmine"]["status"] and parent_id in change_parent_ids and category_name
            target_deliver_item_groups = []
            if is_regular_deliver:
                target_deliver_item_groups.append(deliver_items)
            if is_change_deliver:
                target_deliver_item_groups.append(change_deliver_items)
            if project_name not in CONFIG["redmine"]["excluded_projects"] and target_deliver_item_groups:
                mode = category_name
                reason = []
                deliverables = ""
                design_commit_id = ""
                main_commit_id = ""
                pin_commit_id = ""

                # 获取交付物信息
                for field in issue.custom_fields:
                    if field["name"] in CONFIG["redmine"]["reason_fields"] and field["value"] is not None:
                        reason.append(field["value"])
                    elif field["name"] == CONFIG["redmine"]["custom_fields"]["deliverables"]:
                        deliverables = field["value"]
                    elif field["name"] == CONFIG["redmine"]["custom_fields"]["design_commit_id"]:
                        design_commit_id = field["value"]
                    elif field["name"] == CONFIG["redmine"]["custom_fields"]["main_commit_id"]:
                        main_commit_id = field["value"]
                    elif field["name"] == CONFIG["redmine"]["custom_fields"]["pin_commit_id"]:
                        pin_commit_id = field["value"]

                attachment = {}
                for att in issue.attachments:
                    attachment[att.filename] = att.content_url

                # 根据交付物信息拉取文件
                deliver_list = [item.strip() for item in re.split(r"[,，]", deliverables) if item.strip()]
                mode_key = f"{mode}_{'SW' if any('变更申请' in item or '反馈单' in item or item in ('引脚功能缺失清单', '未实现备注') for item in deliver_list) else 'AE'}"
                for target_deliver_items in target_deliver_item_groups:
                    if project_name not in target_deliver_items:
                        target_deliver_items[project_name] = {mode_key : []}
                    elif mode_key not in target_deliver_items[project_name]:
                        target_deliver_items[project_name][mode_key] = []

                invalid_deliverables = set()
                for item in deliver_list:
                    format_reason = deliverable_format_reason(item, mode, project_name)
                    if format_reason:
                        record_missing(missing_items, project_name, mode, f"交付物清单格式:{item}", format_reason)
                        invalid_deliverables.add(item)
                attachment_report_types = {report_type(filename) for filename in attachment}
                attachment_report_types.discard("")
                pulled_targets = set()
                report_downloaded_types = set()
                if has_test_report_deliverable(deliver_list):
                    changed = download_and_fill_feature_reports(
                        report_packager_config,
                        issue.id,
                        Path(origin_path) / project_name / mode,
                        {"test_report"},
                    )
                    print(f"REPORT_FILLED={issue.id}\t{len(changed)}")
                    for target_deliver_items in target_deliver_item_groups:
                        target_deliver_items[project_name][mode_key].append("测试报告")
                for deliver in deliver_list:
                    if "评审表" in deliver or "反馈单" in deliver:
                        report_types = set()
                        if "评审表" in deliver:
                            report_types.add("review")
                        if "反馈单" in deliver:
                            report_types.add("nonconformity")
                        report_types.update(attachment_report_types & {"review", "nonconformity"})
                        pending_report_types = report_types - report_downloaded_types
                        if pending_report_types:
                            changed = download_and_fill_feature_reports(
                                report_packager_config,
                                issue.id,
                                Path(origin_path) / project_name / mode,
                                pending_report_types,
                            )
                            print(f"REPORT_FILLED={issue.id}\t{len(changed)}")
                            report_downloaded_types.update(pending_report_types)
                        if "review" in report_types:
                            for target_deliver_items in target_deliver_item_groups:
                                target_deliver_items[project_name][mode_key].append("评审表")
                        if "nonconformity" in report_types:
                            for target_deliver_items in target_deliver_item_groups:
                                target_deliver_items[project_name][mode_key].append("反馈单")
                        continue

                    architecture_request = architecture_branch_request(deliver)
                    if architecture_request is not None:
                        if deliver in invalid_deliverables:
                            continue
                        request_mode, branches = architecture_request
                        include_default = request_mode != "only"

                        if include_default and "架构设计" not in pulled_targets:
                            # 默认分支保持原有 commit id 校验规则；额外分支直接拉取各自最新 HEAD。
                            if not git_clone(project_name, mode, design_commit_id, "架构设计", projects, check_commit=is_regular_deliver):
                                print(f"Error: {project_name}下的问题单{issue.id}的架构设计的commit_id不正确")
                            pulled_targets.add("架构设计")

                        if request_mode == "all":
                            default_repo_path = design_repo_path(project_name, mode)
                            default_branch, remote_branches = architecture_remote_branches(default_repo_path)
                            if remote_branches is None:
                                record_missing(
                                    missing_items,
                                    project_name,
                                    mode,
                                    "架构设计@all",
                                    "架构设计要求拉取全部分支，但无法获取远程分支列表",
                                )
                                remote_branches = []
                            branches = [branch for branch in remote_branches if branch != default_branch]

                        branches = list(dict.fromkeys(branches))
                        architecture_items = (["架构设计"] if include_default else []) + [
                            f"架构设计@{branch}" for branch in branches
                        ]
                        if (
                            mode.casefold() in GPIO_MODE_NAMES
                            and not any("pin" in item for item in deliver_list)
                            and "pin_func_list" not in pulled_targets
                        ):
                            git_clone(project_name, mode, "", "pin_func_list", projects, check_commit=False)
                            pulled_targets.add("pin_func_list")
                            architecture_items.append("pin_function_list")
                        for target_deliver_items in target_deliver_item_groups:
                            package_items = target_deliver_items[project_name][mode_key]
                            package_items.extend(item for item in architecture_items if item not in package_items)

                        for branch in branches:
                            branch_target = f"架构设计@{branch}"
                            if branch_target in pulled_targets:
                                continue
                            git_clone(
                                project_name,
                                mode,
                                "",
                                "架构设计",
                                projects,
                                check_commit=False,
                                branch=branch,
                            )
                            pulled_targets.add(branch_target)
                        continue

                    if deliver.upper() in LIBRARY_FOLDERS and deliver in invalid_deliverables:
                        continue
                    for target_deliver_items in target_deliver_item_groups:
                        target_deliver_items[project_name][mode_key].append(deliver)
                    hal_info = hal_deliverable_info(deliver)
                    requirement_keyword = series_requirement_keyword(deliver)
                    if ".ftl" in deliver and "ftl" not in pulled_targets:
                        if not git_clone(project_name, mode, main_commit_id, "ftl", projects, check_commit=is_regular_deliver):
                            print(f"Error: {project_name}下的问题单{issue.id}的main&libopt的commit_id不正确")
                        pulled_targets.add("ftl")
                    elif "pin" in deliver and "pin_func_list" not in pulled_targets:
                        if not git_clone(project_name, mode, pin_commit_id, "pin_func_list", projects, check_commit=is_regular_deliver):
                            print(f"Error: {project_name}下的问题单{issue.id}的pin_function_list的commit_id不正确")
                        pulled_targets.add("pin_func_list")
                    elif "PeriConfig.xml" in deliver and "PeriConfig.xml" not in pulled_targets:
                        git_clone(project_name, mode, "", "PeriConfig.xml", projects)
                        pulled_targets.add("PeriConfig.xml")
                    elif deliver.casefold() == PERIPHERAL_XML_FIELD and PERIPHERAL_XML_FIELD not in pulled_targets:
                        git_clone(project_name, mode, "", PERIPHERAL_XML_FIELD, projects, check_commit=False)
                        pulled_targets.add(PERIPHERAL_XML_FIELD)
                    elif "tips" in deliver.lower() and "tips" not in pulled_targets:
                        git_clone(project_name, mode, "", "tips", projects)
                        pulled_targets.add("tips")
                    elif re.fullmatch(rf"{re.escape(code_series)}_E[A-Z0-9]+", deliver.strip(), re.IGNORECASE) and "example" not in pulled_targets:
                        git_clone(project_name, mode, "", "example", projects)
                        pulled_targets.add("example")
                    elif re.fullmatch(rf"{re.escape(code_series)}_C[A-Z0-9]+", deliver.strip(), re.IGNORECASE) and "case" not in pulled_targets:
                        git_clone(project_name, mode, "", "case", projects)
                        pulled_targets.add("case")
                    elif hal_info and hal_info[1] == "C" and "hal_c" not in pulled_targets:
                        git_clone(project_name, mode, "", "hal_c", projects)
                        pulled_targets.add("hal_c")
                    elif hal_info and hal_info[1] == "H" and "hal_h" not in pulled_targets:
                        git_clone(project_name, mode, "", "hal_h", projects)
                        pulled_targets.add("hal_h")
                    elif deliver.upper() in LIBRARY_FOLDERS and "hal_l" not in pulled_targets:
                        git_clone(project_name, mode, "", "hal_l", projects, check_commit=False)
                        pulled_targets.add("hal_l")
                    elif requirement_keyword and requirement_keyword not in pulled_targets:
                        get_peripheral_feature_analysis_form(
                            dingtalk_series,
                            mode,
                            os.path.join(origin_path, project_name, mode),
                            target_keyword=requirement_keyword,
                        )
                        pulled_targets.add(requirement_keyword)
                    elif deliver == "引脚功能缺失清单" and "引脚功能缺失清单" not in pulled_targets:
                        get_peripheral_feature_analysis_form(
                            dingtalk_series,
                            mode,
                            os.path.join(origin_path, project_name, mode),
                            target_keyword="引脚功能缺失清单",
                        )
                        pulled_targets.add("引脚功能缺失清单")
                    elif "功能分析表" in deliver and "外设功能分析表" not in pulled_targets:
                        get_peripheral_feature_analysis_form(dingtalk_series, mode, os.path.join(origin_path, project_name, mode))
                        pulled_targets.add("外设功能分析表")
                    elif "变更申请" in deliver:
                        get_change_application_form(dingtalk_series, mode, os.path.join(origin_path, project_name, mode))
                        if change_application_form:
                            version_error = change_application_version_error(deliver, change_application_form)
                            if version_error:
                                record_missing(missing_items, project_name, mode, "变更申请单版本", version_error)
                for target_deliver_items in target_deliver_item_groups:
                    target_deliver_items[project_name][mode_key] = list(set(target_deliver_items[project_name][mode_key]))

        except Exception as e:
            record_missing(missing_items, getattr(getattr(issue, "project", None), "name", ""),
                           str(issue.id), "问题单处理", str(e))

    for project_name in need_project_examples:
        mapped_example_repo = example_mapping_table.get(project_name.lower(), "")
        if mapped_example_repo:
            example_group = CONFIG["repositories"]["hal_group"]
        else:
            series = CONFIG["repositories"]["series_aliases"].get(project_name.casefold(), project_name.lower())
            example_group = CONFIG["repositories"]["architecture_group"].format(series=series)
        if example_group not in projects:
            print(f"Error: {project_name}找不到example仓库分组{example_group}")
            continue
        for url in projects[example_group]:
            repo_name = url.split("/")[-1].replace(".git", "")
            if mapped_example_repo:
                if repo_name.casefold() != mapped_example_repo.casefold():
                    continue
            elif "example" not in repo_name.lower():
                continue
            if project_name not in deliver_items:
                deliver_items[project_name] = {}
            if repo_name not in deliver_items[project_name]:
                deliver_items[project_name][repo_name] = []
            if "EXAMPLE" not in deliver_items[project_name][repo_name]:
                deliver_items[project_name][repo_name].append("EXAMPLE")
            os.makedirs(os.path.join(origin_path, project_name, repo_name), exist_ok=True)
            os.chdir(os.path.join(origin_path, project_name, repo_name))
            if not os.path.exists(os.path.join(origin_path, project_name, repo_name, repo_name)):
                subprocess.run(["git", "clone", url], check=True)

    for project_name in need_project_designs:
        series = CONFIG["repositories"]["series_aliases"].get(project_name.casefold(), project_name.lower())
        design_group = CONFIG["repositories"]["architecture_group"].format(series=series)
        if design_group not in projects:
            print(f"Error: {project_name}找不到架构设计仓库分组{design_group}")
            continue
        gpio_mode = ""
        for url in projects[design_group]:
            mode = url.split("/")[-1].replace(".git", "")
            if not mode or mode.lower() == "common" or "example" in mode.lower() or mode.casefold().endswith("_exfunction"):
                continue
            if mode.casefold() in GPIO_MODE_NAMES:
                gpio_mode = mode
            mode_key = f"{mode}_AE"
            if project_name not in deliver_items:
                deliver_items[project_name] = {}
            if mode_key not in deliver_items[project_name]:
                deliver_items[project_name][mode_key] = []
            if "架构设计" not in deliver_items[project_name][mode_key]:
                deliver_items[project_name][mode_key].append("架构设计")
            # 全量架构设计仍按仓库默认分支拉取最新版本。
            git_clone(project_name, mode, "", "架构设计", projects)
            reference_commit = design_reference_commits.get(project_name.casefold(), {}).get(mode.casefold(), "")
            if reference_commit and not git_clone(
                project_name,
                mode,
                "",
                "架构设计",
                projects,
                check_commit=False,
                reference_commit=reference_commit,
            ):
                print(f"Error: {project_name}的{mode}设计参考git commit id不正确")
        if gpio_mode:
            # 全量架构设计的 GPIO 目录需要合并 Common 仓库中的 pin_function_list 文件。
            if git_clone(project_name, gpio_mode, "", "pin_func_list", projects, check_commit=False):
                copy_directory_contents(os.getcwd(), design_repo_path(project_name, gpio_mode))

    for project_name in need_project_firmwares:
        hal_repo_name = hal_mapping_table.get(project_name.lower(), "")
        if not hal_repo_name:
            print(f"Error: {project_name}找不到Firmware对应HAL仓库映射")
            continue
        if CONFIG["repositories"]["hal_group"] not in projects:
            print(f"Error: {project_name}找不到配置的HAL仓库分组")
            continue

        for url in projects[CONFIG["repositories"]["hal_group"]]:
            repo_name = url.split("/")[-1].replace(".git", "")
            if repo_name.lower() != hal_repo_name.lower():
                continue
            firmware_path = os.path.join(origin_path, project_name, "Firmware")
            os.makedirs(firmware_path, exist_ok=True)
            os.chdir(firmware_path)
            if not os.path.exists(os.path.join(firmware_path, repo_name)):
                subprocess.run(["git", "clone", url], check=True)

    package_paths = []
    package_paths.extend(packaged_deliverables_SW(deliver_items, missing_items))
    package_paths.extend(packaged_deliverables_AE(deliver_items, missing_items))
    package_paths.extend(packaged_deliverables_SW(change_deliver_items, missing_items, "Change"))
    package_paths.extend(packaged_deliverables_AE(change_deliver_items, missing_items, "Change"))
    packaged_examples(deliver_items, missing_items)
    packaged_firmwares(need_project_firmwares, missing_items)
    generate_deliver_list_from_packages(package_paths)
    # 清单生成后再去掉 branches 中间层，避免分支目录被误识别成新的外设模块。
    for package_path in package_paths:
        for branches_path in glob.glob(os.path.join(glob.escape(package_path), "*", "branches")):
            for branch_name in os.listdir(branches_path):
                source_path = os.path.join(branches_path, branch_name)
                target_path = os.path.join(os.path.dirname(branches_path), branch_name)
                if os.path.exists(target_path):
                    raise FileExistsError(f"分支目录上移目标已存在: {target_path}")
                shutil.move(source_path, target_path)
                print(f"ARCH_BRANCH_FLATTEN={source_path}\t{target_path}")
            os.rmdir(branches_path)
    print_missing_items(missing_items)
    return 1 if missing_items else 0

if __name__ == "__main__":
    raise SystemExit(main())
